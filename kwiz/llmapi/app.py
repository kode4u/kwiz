"""
LLM API Service for Question Generation
Generates structured MCQ questions using LLM backends
"""
from flask import Flask, request, jsonify, render_template
from flask_cors import CORS
import os
import logging
import time
import threading
import requests
import hashlib
import io
import base64
import re
from dotenv import load_dotenv
from pydantic import BaseModel, Field
from typing import List, Optional
import json

from metrics_logger import (
    log_generation,
    log_hardware_once,
    log_inacon_metrics,
    get_next_iteration_number,
    get_current_iteration_number,
    reset_iteration_number,
    log_evaluation_iteration,
    get_all_iterations_summary,
    get_iteration_details,
    get_evaluation_dashboard_stats,
)

# Persistent Embedding Cache setup
EMBEDDINGS_CACHE_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "embeddings_cache.json")
embedding_cache_lock = threading.Lock()
embedding_cache = {}

def load_embedding_cache():
    global embedding_cache
    if os.path.isfile(EMBEDDINGS_CACHE_FILE):
        try:
            with open(EMBEDDINGS_CACHE_FILE, 'r', encoding='utf-8') as fh:
                embedding_cache = json.load(fh)
            logger.info(f"Loaded {len(embedding_cache)} cached embeddings from {EMBEDDINGS_CACHE_FILE}")
        except Exception as e:
            logger.warning(f"Failed to load embedding cache: {e}")

def save_embedding_cache():
    try:
        os.makedirs(os.path.dirname(EMBEDDINGS_CACHE_FILE), exist_ok=True)
        with open(EMBEDDINGS_CACHE_FILE, 'w', encoding='utf-8') as fh:
            json.dump(embedding_cache, fh)
    except Exception as e:
        logger.warning(f"Failed to save embedding cache: {e}")

def chunk_text(content: str, target_size: int = 500, max_size: int = 800) -> list:
    """Robust text chunker that guarantees chunk size <= max_size even for texts without newlines."""
    content = (content or '').strip()
    if not content:
        return []

    # Normalize line endings
    content = content.replace('\r\n', '\n').replace('\r', '\n')

    # Split into logical blocks: paragraphs, lines, or sentences
    raw_blocks = re.split(r'(\n{2,}|\n|(?<=[.!?])\s+)', content)

    chunks = []
    current_chunk = []
    current_len = 0

    for block in raw_blocks:
        if not block:
            continue

        # If a single block exceeds max_size (e.g. giant unpunctuated string), slice by words
        if len(block) > max_size:
            words = block.split(' ')
            sub_chunk = []
            sub_len = 0
            for w in words:
                if sub_len + len(w) + 1 > target_size and sub_chunk:
                    chunks.append(' '.join(sub_chunk).strip())
                    sub_chunk = []
                    sub_len = 0
                sub_chunk.append(w)
                sub_len += len(w) + 1
            if sub_chunk:
                chunks.append(' '.join(sub_chunk).strip())
            continue

        if current_len + len(block) > max_size and current_chunk:
            combined = ''.join(current_chunk).strip()
            if combined:
                chunks.append(combined)
            current_chunk = []
            current_len = 0

        current_chunk.append(block)
        current_len += len(block)

        if current_len >= target_size:
            combined = ''.join(current_chunk).strip()
            if combined:
                chunks.append(combined)
            current_chunk = []
            current_len = 0

    if current_chunk:
        combined = ''.join(current_chunk).strip()
        if combined:
            chunks.append(combined)

    cleaned = [c for c in chunks if c.strip()]
    return cleaned if cleaned else [content[:max_size]]

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

load_dotenv()
load_embedding_cache()

app = Flask(__name__)
app.config['JSON_AS_ASCII'] = False
try:
    app.json.ensure_ascii = False
except AttributeError:
    pass
CORS(app)

# Configuration
LLM_BACKEND = os.getenv('LLM_BACKEND', 'local')
OPENAI_API_KEY = os.getenv('OPENAI_API_KEY', '')
OPENAI_MODEL = os.getenv('OPENAI_MODEL', 'gpt-4o')
GEMINI_API_KEY = os.getenv('GEMINI_API_KEY', '')
GEMINI_MODEL = os.getenv('GEMINI_MODEL', 'gemini-2.5-flash')
LOCAL_LLM_URL = os.getenv('LOCAL_LLM_URL', 'http://localhost:11434')  # Ollama default
MAX_QUESTIONS = int(os.getenv('MAX_QUESTIONS', '20'))
DEFAULT_LANGUAGE = os.getenv('DEFAULT_LANGUAGE', 'en')
LOCAL_LLM_TIMEOUT = int(os.getenv('LOCAL_LLM_TIMEOUT', '1200'))
MAX_LESSON_CONTEXT_CHARS = int(os.getenv('MAX_LESSON_CONTEXT_CHARS', '12000'))
LOCAL_GEN_BATCH_SIZE = int(os.getenv('LOCAL_GENERATION_BATCH_SIZE', '3'))
WEBHOOK_TIMEOUT = int(os.getenv('WEBHOOK_TIMEOUT', '30'))
# Optional: comma-separated list of Ollama models to pre-pull on startup
OLLAMA_PRELOAD_MODELS = os.getenv('OLLAMA_PRELOAD_MODELS', '').strip()
OLLAMA_MODEL_DEFAULT = os.getenv('OLLAMA_MODEL', 'qwen2.5-coder:7b')

log_hardware_once(LOCAL_LLM_URL, OLLAMA_MODEL_DEFAULT, LLM_BACKEND)


class QuestionRequest(BaseModel):
    topic: str = Field(..., description="Topic for question generation")
    level: str = Field(default="medium", description="Difficulty level: easy, medium, hard")
    n_questions: int = Field(default=1, ge=1, le=MAX_QUESTIONS, description="Number of questions")
    language: str = Field(default=DEFAULT_LANGUAGE, description="Language code: en, km")
    bloom_level: Optional[str] = Field(default=None, description="Bloom's taxonomy level")
    context: Optional[str] = Field(default=None, description="Optional lesson material to base questions on")
    backend: Optional[str] = Field(default=None, description="LLM backend: openai, gemini, local")
    model: Optional[str] = Field(default=None, description="Override model name for local backend")
    openai_api_key: Optional[str] = Field(default=None, description="Per-request OpenAI API key override")
    gemini_api_key: Optional[str] = Field(default=None, description="Per-request Gemini API key override")
    learning_outcomes: Optional[str] = Field(default=None, description="Optional target learning outcomes")
    category_name: Optional[str] = Field(default="", description="Category label")
    top_k: Optional[int] = Field(default=3, ge=1, le=10, description="Top-K retrieved chunks")
    max_context_chars: Optional[int] = Field(default=None, description="Max context length limit")
    enable_incremental_cache: Optional[bool] = Field(default=True, description="Enable SHA-256 chunk caching")
    pipeline_mode: Optional[str] = Field(default="proposed", description="Pipeline mode: proposed (or INACON), full_reindex, static_context, no_rag")
    corpus_tokens: Optional[int] = Field(default=0, description="Corpus scale tokens for evaluation tracking")
    change_ratio: Optional[float] = Field(default=0.0, description="Corpus change ratio: 0.0 to 1.0")
    iteration_number: Optional[int] = Field(default=None, description="Explicit evaluation iteration number")
    request_uuid: Optional[str] = Field(default=None, description="Unique tracking UUID")
    question_type: Optional[str] = Field(default="code", description="Question modality: code, mixed, conceptual")


class Choice(BaseModel):
    text: str
    is_correct: bool


class Question(BaseModel):
    question: str
    choices: List[Choice]
    correct_index: int
    difficulty: str
    bloom_level: str
    explanation: Optional[str] = None


class QuestionResponse(BaseModel):
    questions: List[Question]
    metadata: dict


class AsyncGenerateRequest(QuestionRequest):
    request_uuid: str = Field(..., description="Job id shared with Moodle")
    webhook_url: str = Field(..., description="Moodle callback URL when generation finishes")
    webhook_token: str = Field(..., description="Shared secret for webhook auth")


def _preload_ollama_models():
    """
    Optionally pre-pull a set of Ollama models on startup.
    Controlled via OLLAMA_PRELOAD_MODELS env var (comma-separated list).
    """
    if not OLLAMA_PRELOAD_MODELS:
        return
    if not LOCAL_LLM_URL:
        logger.warning("OLLAMA_PRELOAD_MODELS is set but LOCAL_LLM_URL is empty; skipping preload.")
        return

    models = [m.strip() for m in OLLAMA_PRELOAD_MODELS.split(',') if m.strip()]
    if not models:
        return

    try:
        logger.info(f"Preloading Ollama models: {models}")

        # Get currently available models
        try:
            resp = requests.get(f"{LOCAL_LLM_URL}/api/tags", timeout=15)
            resp.raise_for_status()
            tags = resp.json().get('models', []) or []
            existing = {m.get('name') or m.get('model') for m in tags}
        except Exception as e:
            logger.warning(f"Could not list existing Ollama models at {LOCAL_LLM_URL}: {e}")
            existing = set()

        # Pull any missing models
        for model in models:
            if model in existing:
                logger.info(f"Ollama model already present: {model}")
                continue
            try:
                logger.info(f"Pre-pulling Ollama model: {model}")
                pull_resp = requests.post(
                    f"{LOCAL_LLM_URL}/api/pull",
                    json={"model": model, "stream": False},
                    timeout=1800,  # up to 30 minutes for large models
                )
                if pull_resp.status_code != 200:
                    logger.error(
                        "Failed to pull Ollama model %s: %s %s",
                        model,
                        pull_resp.status_code,
                        pull_resp.text[:200],
                    )
                else:
                    logger.info(f"Successfully pulled Ollama model: {model}")
            except Exception as e:
                logger.error(f"Error while pulling Ollama model {model}: {e}")
    except Exception as e:
        logger.error(f"Error during Ollama preload: {e}", exc_info=True)


def format_lesson_context(context: Optional[str]) -> str:
    """Format optional pasted lesson text for inclusion in generation prompts."""
    if not context or not str(context).strip():
        return ''
    text = str(context).strip()
    if len(text) > MAX_LESSON_CONTEXT_CHARS:
        text = text[:MAX_LESSON_CONTEXT_CHARS] + "\n[... lesson truncated for length ...]"
    return (
        "\n- Base questions primarily on the following lesson material "
        "(stay faithful to the content; do not invent facts beyond it):\n"
        "---\n"
        f"{text}\n"
        "---\n"
    )


def format_question_type(question_type: Optional[str]) -> str:
    """Format modality instructions for programming questions."""
    qtype = (question_type or 'code').lower().strip()
    if 'conceptual' in qtype or qtype == 'theory':
        return (
            "- Question Modality: CONCEPTUAL & THEORETICAL.\n"
            "  * Test deep understanding of programming concepts, terminology, syntax rules, algorithms, and data structure characteristics without code snippets.\n"
            "  * Do not include code blocks in the question text.\n"
        )
    elif 'mixed' in qtype:
        return (
            "- Question Modality: BALANCED MIX (Code-Centric and Conceptual).\n"
            "  * Provide a balanced mix: questions with executable code snippets (```python ... ```) testing tracing/output prediction, and questions testing conceptual understanding.\n"
        )
    else:
        return (
            "- Question Modality: STRICTLY CODE-CENTRIC (Execution Tracing & Output Analysis).\n"
            "  * Every single question MUST include an executable code snippet (3 to 10 lines) inside a ```python ... ``` block in the question text.\n"
            "  * The question must require students to trace the code, predict the exact printed output, determine variable states, or find syntax/runtime bugs.\n"
            "  * Ensure all code snippets are 100% syntactically valid Python 3 that can compile without SyntaxError.\n"
            "  * Distractors (incorrect choices) must reflect realistic cognitive tracing mistakes (such as off-by-one errors, 0-indexing confusion, incorrect operator precedence, mutable reference misunderstandings).\n"
            "  * DO NOT generate generic definitions or theory-only questions without code blocks.\n"
        )


def format_learning_outcomes(learning_outcomes: Optional[str]) -> str:
    """Format target learning outcomes for inclusion in generation prompts."""
    if not learning_outcomes or not str(learning_outcomes).strip():
        return ''
    return f"- Target Learning Outcomes: {learning_outcomes}\n"


def strip_json_fences(content: str) -> str:
    """Safely strip outer markdown code fences (e.g. ```json ... ```) without affecting inner code blocks."""
    if not content:
        return ''
    text = content.strip()
    if text.startswith('```'):
        text = re.sub(r'^```(?:json|JSON)?\s*\n?', '', text)
        text = re.sub(r'\n?```\s*$', '', text)
    return text.strip()


def parse_llm_json(content: str) -> List[dict]:
    """Robustly parse JSON response from LLM, handling markdown fences, extra text, and trailing commas."""
    cleaned = strip_json_fences(content)
    
    # Attempt 1: Direct JSON decode
    try:
        data = json.loads(cleaned)
        if isinstance(data, dict) and 'questions' in data and isinstance(data['questions'], list):
            return data['questions']
        if isinstance(data, dict):
            return [data]
        if isinstance(data, list):
            return data
    except json.JSONDecodeError:
        pass
        
    # Attempt 2: Extract outermost matching bracket/brace
    start_bracket = cleaned.find('[')
    start_brace = cleaned.find('{')
    start_idx = -1
    if start_bracket != -1 and start_brace != -1:
        start_idx = min(start_bracket, start_brace)
    elif start_bracket != -1:
        start_idx = start_bracket
    elif start_brace != -1:
        start_idx = start_brace
        
    if start_idx != -1:
        end_idx = max(cleaned.rfind(']'), cleaned.rfind('}'))
        if end_idx > start_idx:
            sub = cleaned[start_idx:end_idx + 1]
            try:
                data = json.loads(sub)
                if isinstance(data, dict) and 'questions' in data and isinstance(data['questions'], list):
                    return data['questions']
                if isinstance(data, dict):
                    return [data]
                if isinstance(data, list):
                    return data
            except json.JSONDecodeError:
                # Attempt 3: Strip trailing commas before closing braces/brackets
                sub_fixed = re.sub(r',\s*([\]}])', r'\1', sub)
                try:
                    data = json.loads(sub_fixed)
                    if isinstance(data, dict) and 'questions' in data and isinstance(data['questions'], list):
                        return data['questions']
                    if isinstance(data, dict):
                        return [data]
                    if isinstance(data, list):
                        return data
                except json.JSONDecodeError:
                    pass

    preview = cleaned[:250] if cleaned else '(empty response from model)'
    raise Exception(f"Invalid JSON from LLM: Expecting value. Response preview: {preview}")


import math

def dot_product(v1, v2):
    return sum(x * y for x, y in zip(v1, v2))

def magnitude(v):
    return math.sqrt(sum(x * x for x in v))

def cosine_similarity(v1, v2):
    mag1 = magnitude(v1)
    mag2 = magnitude(v2)
    if mag1 == 0 or mag2 == 0:
        return 0.0
    return dot_product(v1, v2) / (mag1 * mag2)

def retrieve_relevant_context(
    query: str,
    document_text: str,
    backend: str,
    api_key_override: Optional[str] = None,
    local_model: Optional[str] = None,
    top_k: int = 3,
    max_context_chars: Optional[int] = None,
    enable_cache: bool = True,
    pipeline_mode: str = "proposed"
) -> tuple[str, dict]:
    """Perform RAG vector retrieval with decomposed timing instrumentation."""
    metrics = {
        't_extract_ms': 0.0,
        't_chunk_ms': 0.0,
        't_hash_ms': 0.0,
        't_lookup_ms': 0.0,
        't_embed_ms': 0.0,
        't_index_ms': 0.0,
        't_kb_ms': 0.0,
        't_query_embed_ms': 0.0,
        't_retrieval_ms': 0.0,
        'cache_hits': 0,
        'cache_misses': 0,
        'hit_ratio': 0.0,
        'chunks_count': 0,
    }

    if not document_text or not document_text.strip() or pipeline_mode == "no_rag":
        return "", metrics

    # Baseline/Ablation C: Static Context Window (No Retrieval)
    if pipeline_mode == "static_context":
        char_limit = max_context_chars or 6000
        return document_text[:char_limit], metrics

    # 1. T_extract: Extract and normalize text
    t_ext_start = time.perf_counter()
    clean_text = document_text.strip()
    metrics['t_extract_ms'] = (time.perf_counter() - t_ext_start) * 1000.0

    # Short document bypass
    if len(clean_text) <= 1500 and pipeline_mode != "full_reindex":
        metrics['t_kb_ms'] = metrics['t_extract_ms']
        return clean_text, metrics

    # 2. T_chunk: Semantic / sliding window chunking
    t_chunk_start = time.perf_counter()
    chunks = chunk_text(clean_text)
    metrics['t_chunk_ms'] = (time.perf_counter() - t_chunk_start) * 1000.0
    metrics['chunks_count'] = len(chunks)

    if not chunks:
        metrics['t_kb_ms'] = metrics['t_extract_ms'] + metrics['t_chunk_ms']
        return clean_text[:6000], metrics

    # Model resolution for embeddings
    if backend == 'openai':
        model_name = "text-embedding-3-small"
    elif backend == 'gemini':
        model_name = "models/text-embedding-004"
    elif backend == 'local':
        model_name = os.getenv('OLLAMA_EMBED_MODEL', 'nomic-embed-text') or local_model or os.getenv('OLLAMA_MODEL', 'qwen2.5-coder:7b')
    else:
        model_name = "unknown"

    def compute_single_embedding(text: str) -> list[float]:
        try:
            if backend == 'openai':
                from openai import OpenAI
                client = OpenAI(api_key=api_key_override or OPENAI_API_KEY)
                resp = client.embeddings.create(input=[text], model=model_name)
                return resp.data[0].embedding
            elif backend == 'gemini':
                import google.generativeai as genai
                genai.configure(api_key=api_key_override or GEMINI_API_KEY)
                resp = genai.embed_content(model=model_name, content=text)
                return resp['embedding']
            elif backend == 'local':
                resp = requests.post(f"{LOCAL_LLM_URL}/api/embeddings", json={"model": model_name, "prompt": text}, timeout=60)
                resp.raise_for_status()
                return resp.json()['embedding']
        except Exception as embed_err:
            logger.warning(f"Failed to generate embedding for text with '{model_name}': {embed_err}")
            return []
        return []

    # 3. T_hash, T_lookup, T_embed, T_index for knowledge base chunks
    chunk_vectors = []
    valid_chunks = []
    new_cache_entries = {}

    for c in chunks:
        # T_hash
        t_h_start = time.perf_counter()
        hash_key = hashlib.sha256(f"{model_name}:{c}".encode('utf-8')).hexdigest()
        metrics['t_hash_ms'] += (time.perf_counter() - t_h_start) * 1000.0

        cached_vec = None
        if enable_cache and pipeline_mode != "full_reindex":
            # T_lookup
            t_lk_start = time.perf_counter()
            with embedding_cache_lock:
                cached_vec = embedding_cache.get(hash_key)
            metrics['t_lookup_ms'] += (time.perf_counter() - t_lk_start) * 1000.0

        if cached_vec is not None:
            metrics['cache_hits'] += 1
            chunk_vectors.append(cached_vec)
            valid_chunks.append(c)
        else:
            metrics['cache_misses'] += 1
            # T_embed
            t_emb_start = time.perf_counter()
            vec = compute_single_embedding(c)
            metrics['t_embed_ms'] += (time.perf_counter() - t_emb_start) * 1000.0
            if vec:
                chunk_vectors.append(vec)
                valid_chunks.append(c)
                if enable_cache and pipeline_mode != "full_reindex":
                    new_cache_entries[hash_key] = vec

    # T_index: persist new embeddings into cache
    if new_cache_entries:
        t_idx_start = time.perf_counter()
        with embedding_cache_lock:
            embedding_cache.update(new_cache_entries)
            save_embedding_cache()
        metrics['t_index_ms'] = (time.perf_counter() - t_idx_start) * 1000.0

    total_chunks = metrics['cache_hits'] + metrics['cache_misses']
    metrics['hit_ratio'] = metrics['cache_hits'] / total_chunks if total_chunks > 0 else 0.0
    metrics['t_kb_ms'] = (
        metrics['t_extract_ms']
        + metrics['t_chunk_ms']
        + metrics['t_hash_ms']
        + metrics['t_lookup_ms']
        + metrics['t_embed_ms']
        + metrics['t_index_ms']
    )

    if not chunk_vectors:
        return clean_text[:6000], metrics

    # 4. T_query_embed: Compute embedding for the incoming query
    t_qe_start = time.perf_counter()
    query_vector = None
    query_hash = hashlib.sha256(f"{model_name}:query:{query}".encode('utf-8')).hexdigest()
    if enable_cache and pipeline_mode != "full_reindex":
        with embedding_cache_lock:
            query_vector = embedding_cache.get(query_hash)
    if query_vector is None:
        query_vector = compute_single_embedding(query)
        if query_vector and enable_cache and pipeline_mode != "full_reindex":
            with embedding_cache_lock:
                embedding_cache[query_hash] = query_vector
                save_embedding_cache()
    metrics['t_query_embed_ms'] = (time.perf_counter() - t_qe_start) * 1000.0

    if not query_vector:
        return clean_text[:6000], metrics

    # 5. T_retrieval: Cosine similarity computation, ranking, and Top-K extraction
    t_ret_start = time.perf_counter()
    similarities = [cosine_similarity(query_vector, cv) for cv in chunk_vectors]
    ranked_indices = sorted(range(len(similarities)), key=lambda i: similarities[i], reverse=True)
    k = max(1, min(int(top_k), len(ranked_indices)))
    relevant = [valid_chunks[idx] for idx in ranked_indices[:k]]
    retrieved_text = "\n\n---\n\n".join(relevant)
    if max_context_chars and len(retrieved_text) > max_context_chars:
        retrieved_text = retrieved_text[:max_context_chars]
    metrics['t_retrieval_ms'] = (time.perf_counter() - t_ret_start) * 1000.0

    return retrieved_text, metrics


def generate_with_openai(topic: str, level: str, n_questions: int, language: str, bloom_level: Optional[str], context: Optional[str], api_key_override: Optional[str] = None, learning_outcomes: Optional[str] = None, question_type: Optional[str] = "code") -> List[Question]:
    """Generate questions using OpenAI API"""
    try:
        from openai import OpenAI
        
        # Initialize client with just the API key
        # OpenAI library 2.x+ uses simple initialization
        api_key = api_key_override or OPENAI_API_KEY
        client = OpenAI(api_key=api_key)
        
        is_code = 'code' in (question_type or 'code').lower() and 'conceptual' not in (question_type or 'code').lower()
        q_example = (
            "What is the output of the following code snippet?\\n```python\\nx = 10\\nfor i in range(3):\\n    x += i\\nprint(x)\\n```"
            if is_code
            else "Which of the following statements is true regarding Python loops?"
        )

        prompt = f"""Generate a JSON array of exactly {n_questions} multiple-choice question(s) on the topic: "{topic}"

Requirements:
- Difficulty level: {level}
- Language: {language}
- Bloom's taxonomy level: {bloom_level or 'comprehension'}
{format_question_type(question_type)}
{format_learning_outcomes(learning_outcomes)}
{format_lesson_context(context)}

For each question, provide:
1. Clear question text (for code-centric questions, the executable Python snippet MUST be embedded inside ```python ... ``` within the question string)
2. Exactly 4 answer choices (only one correct)
3. The index (0-3) of the correct answer
4. A brief explanation

CRITICAL INSTRUCTIONS:
- Return ONLY a valid JSON array. Do not wrap the JSON output in markdown code fences or conversational text.
- Escape all internal quotation marks and newlines inside JSON strings properly.

Format as JSON array:
[
  {{
    "question": "{q_example}",
    "choices": [
      {{"text": "Choice 1", "is_correct": true}},
      {{"text": "Choice 2", "is_correct": false}},
      {{"text": "Choice 3", "is_correct": false}},
      {{"text": "Choice 4", "is_correct": false}}
    ],
    "correct_index": 0,
    "difficulty": "{level}",
    "bloom_level": "{bloom_level or 'comprehension'}",
    "explanation": "Brief explanation"
  }}
]"""

        response = client.chat.completions.create(
            model=OPENAI_MODEL,
            messages=[
                {"role": "system", "content": "You are an expert educational content generator. Always return valid JSON."},
                {"role": "user", "content": prompt}
            ],
            temperature=0.7,
            max_tokens=2000
        )
        
        content = response.choices[0].message.content or ""
        questions_data = parse_llm_json(content)
        questions = []
        
        for q_data in questions_data:
            correct_index = None
            choices_list = []
            for idx, choice_data in enumerate(q_data['choices']):
                choices_list.append(Choice(**choice_data))
                if choice_data.get('is_correct'):
                    correct_index = idx
            
            if correct_index is None:
                correct_index = q_data.get('correct_index', 0)
            
            questions.append(Question(
                question=q_data['question'],
                choices=choices_list,
                correct_index=correct_index,
                difficulty=q_data.get('difficulty', level),
                bloom_level=q_data.get('bloom_level', bloom_level or 'comprehension'),
                explanation=q_data.get('explanation')
            ))
        
        return questions
        
    except Exception as e:
        raise Exception(f"OpenAI generation error: {str(e)}")


def generate_with_gemini(topic: str, level: str, n_questions: int, language: str, bloom_level: Optional[str], context: Optional[str], api_key_override: Optional[str] = None, learning_outcomes: Optional[str] = None, question_type: Optional[str] = "code") -> List[Question]:
    """Generate questions using Google Gemini API"""
    try:
        import google.generativeai as genai
        
        effective_api_key = api_key_override or GEMINI_API_KEY
        if not effective_api_key:
            raise Exception("Gemini API key not configured")
        
        genai.configure(api_key=effective_api_key)
        model = genai.GenerativeModel(GEMINI_MODEL)
        
        is_code = 'code' in (question_type or 'code').lower() and 'conceptual' not in (question_type or 'code').lower()
        q_example = (
            "What is the output of the following code snippet?\\n```python\\nx = 10\\nfor i in range(3):\\n    x += i\\nprint(x)\\n```"
            if is_code
            else "Which of the following statements is true regarding Python loops?"
        )

        prompt = f"""Generate a JSON array of exactly {n_questions} multiple-choice question(s) on the topic: "{topic}"

Requirements:
- Difficulty level: {level}
- Language: {language}
- Bloom's taxonomy level: {bloom_level or 'comprehension'}
{format_question_type(question_type)}
{format_learning_outcomes(learning_outcomes)}
{format_lesson_context(context)}

For each question, provide:
1. Clear question text (for code-centric questions, the executable Python snippet MUST be embedded inside ```python ... ``` within the question string)
2. Exactly 4 answer choices (only one correct)
3. The index (0-3) of the correct answer
4. A brief explanation

CRITICAL INSTRUCTIONS:
- Return ONLY a valid JSON array. Do not wrap the JSON output in markdown code fences or conversational text.
- Escape all internal quotation marks and newlines inside JSON strings properly.

Format as JSON array:
[
  {{
    "question": "{q_example}",
    "choices": [
      {{"text": "Choice 1", "is_correct": true}},
      {{"text": "Choice 2", "is_correct": false}},
      {{"text": "Choice 3", "is_correct": false}},
      {{"text": "Choice 4", "is_correct": false}}
    ],
    "correct_index": 0,
    "difficulty": "{level}",
    "bloom_level": "{bloom_level or 'comprehension'}",
    "explanation": "Brief explanation"
  }}
]"""
        
        response = model.generate_content(prompt)
        content = response.text or ""
        questions_data = parse_llm_json(content)
        questions = []
        
        for q_data in questions_data:
            choices = [Choice(text=choice['text'], is_correct=choice['is_correct']) 
                      for choice in q_data['choices']]
            question = Question(
                question=q_data['question'],
                choices=choices,
                correct_index=q_data['correct_index'],
                difficulty=q_data.get('difficulty', level),
                bloom_level=q_data.get('bloom_level', bloom_level or 'comprehension'),
                explanation=q_data.get('explanation')
            )
            questions.append(question)
        
        return questions
        
    except Exception as e:
        raise Exception(f"Gemini generation error: {str(e)}")


def generate_with_local_llm(topic: str, level: str, n_questions: int, language: str,
                            bloom_level: Optional[str], context: Optional[str],
                            model: Optional[str] = None, learning_outcomes: Optional[str] = None,
                            question_type: Optional[str] = "code") -> List[Question]:
    """Generate questions using local LLM (Ollama)"""
    try:
        ollama_model = model or os.getenv('OLLAMA_MODEL', 'qwen2.5-coder:7b')
        logger.info(f"Connecting to Ollama at {LOCAL_LLM_URL} with model {ollama_model}")
        
        is_code = 'code' in (question_type or 'code').lower() and 'conceptual' not in (question_type or 'code').lower()
        q_example = (
            "What is the output of the following code snippet?\\n```python\\nx = 10\\nfor i in range(3):\\n    x += i\\nprint(x)\\n```"
            if is_code
            else "Which of the following statements is true regarding Python loops?"
        )

        prompt = f"""Generate a JSON array of exactly {n_questions} multiple-choice question(s) on the topic: "{topic}"

Requirements:
- Difficulty level: {level}
- Language: {language}
- Bloom's taxonomy level: {bloom_level or 'comprehension'}
{format_question_type(question_type)}
{format_learning_outcomes(learning_outcomes)}
{format_lesson_context(context)}

For each question, provide:
1. Clear question text (for code-centric questions, the executable Python snippet MUST be embedded inside ```python ... ``` within the question string)
2. Exactly 4 answer choices (only one correct)
3. The index (0-3) of the correct answer
4. A brief explanation

CRITICAL INSTRUCTIONS:
- Return ONLY a valid JSON array. Do not wrap the JSON output in markdown code fences or conversational text.
- Escape all internal quotation marks and newlines inside JSON strings properly.

Format as JSON array:
[
  {{
    "question": "{q_example}",
    "choices": [
      {{"text": "Choice 1", "is_correct": true}},
      {{"text": "Choice 2", "is_correct": false}},
      {{"text": "Choice 3", "is_correct": false}},
      {{"text": "Choice 4", "is_correct": false}}
    ],
    "correct_index": 0,
    "difficulty": "{level}",
    "bloom_level": "{bloom_level or 'comprehension'}",
    "explanation": "Brief explanation"
  }}
]"""
        
        # Use Ollama API
        response = requests.post(
            f"{LOCAL_LLM_URL}/api/generate",
            json={
                "model": ollama_model,
                "prompt": prompt,
                "stream": False,
                "options": {
                    "temperature": 0.2,
                    "top_p": 0.9,
                }
            },
            timeout=LOCAL_LLM_TIMEOUT
        )
        
        if response.status_code != 200:
            error_detail = response.text if hasattr(response, 'text') else 'Unknown error'
            logger.error(f"Ollama API error: {response.status_code} - {error_detail}")
            raise Exception(f"Local LLM API error: {response.status_code} - {error_detail}")
        
        logger.info("Successfully received response from Ollama")
        
        result = response.json()
        content = result.get('response', '').strip()
        questions_data = parse_llm_json(content)
        
        # Handle both single object and array responses
        if not isinstance(questions_data, list):
            questions_data = [questions_data]
        
        questions = []
        
        for q_data in questions_data:
            choices = [Choice(text=choice['text'], is_correct=choice['is_correct']) 
                      for choice in q_data['choices']]
            question = Question(
                question=q_data['question'],
                choices=choices,
                correct_index=q_data['correct_index'],
                difficulty=q_data.get('difficulty', level),
                bloom_level=q_data.get('bloom_level', bloom_level or 'comprehension'),
                explanation=q_data.get('explanation')
            )
            questions.append(question)
        
        logger.info(f"Successfully generated {len(questions)} questions")
        return questions
        
    except requests.exceptions.ConnectionError as e:
        error_msg = f"Cannot connect to Ollama at {LOCAL_LLM_URL}. Make sure Ollama is running and accessible from Docker container."
        logger.error(error_msg)
        raise Exception(f"Local LLM generation error: {error_msg} - {str(e)}")
    except requests.exceptions.Timeout as e:
        error_msg = f"Ollama request timed out after {LOCAL_LLM_TIMEOUT} seconds"
        logger.error(error_msg)
        raise Exception(f"Local LLM generation error: {error_msg}")
    except Exception as e:
        logger.error(f"Local LLM generation error: {str(e)}", exc_info=True)
        raise Exception(f"Local LLM generation error: {str(e)}")


@app.route('/health', methods=['GET'])
def health():
    """Health check endpoint"""
    return jsonify({
        'status': 'healthy',
        'backend': LLM_BACKEND,
        'service': 'llmapi'
    }), 200


@app.route('/models/ollama', methods=['GET'])
def list_ollama_models():
    """List models available in local Ollama."""
    try:
        resp = requests.get(f"{LOCAL_LLM_URL}/api/tags", timeout=10)
        resp.raise_for_status()
        return jsonify(resp.json()), 200
    except Exception as e:
        logger.error(f"Error listing Ollama models: {e}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/models/ollama/pull', methods=['POST'])
def pull_ollama_model():
    """Trigger download (pull) of a specific Ollama model on demand."""
    try:
        data = request.get_json(force=True, silent=True) or {}
        model = data.get('model')
        stream = bool(data.get('stream', False))

        if not model:
            return jsonify({'error': "Missing 'model' in request body"}), 400

        logger.info(f"Pulling Ollama model on demand: {model} (stream={stream})")
        resp = requests.post(
            f"{LOCAL_LLM_URL}/api/pull",
            json={"model": model, "stream": stream},
            timeout=1800,  # up to 30 minutes
            stream=stream,
        )

        # If streaming, proxy chunks back to client.
        if stream:
            def generate():
                try:
                    for chunk in resp.iter_content(chunk_size=None):
                        if chunk:
                            yield chunk
                except Exception as e:
                    logger.error(f"Error streaming Ollama pull for {model}: {e}", exc_info=True)
            return app.response_class(generate(), status=resp.status_code, mimetype='application/x-ndjson')

        # Non-streaming: just return JSON/text
        try:
            resp.raise_for_status()
        except Exception:
            logger.error(
                "Failed to pull Ollama model %s: %s %s",
                model,
                resp.status_code,
                getattr(resp, 'text', '')[:200],
            )
            return jsonify({'error': f"Failed to pull model {model}", 'status': resp.status_code, 'detail': resp.text}), resp.status_code

        # Try to pass through JSON if possible
        try:
            return jsonify(resp.json()), resp.status_code
        except Exception:
            return jsonify({'status': 'ok', 'detail': resp.text}), resp.status_code

    except Exception as e:
        logger.error(f"Error pulling Ollama model: {e}", exc_info=True)
        return jsonify({'error': str(e)}), 500


def execute_generation(req: QuestionRequest) -> List[Question]:
    """Run generation for one request (batched for slow local models)."""
    import uuid

    # Ensure unique iteration number and request UUID
    if not getattr(req, 'iteration_number', None):
        req.iteration_number = get_next_iteration_number()
    if not getattr(req, 'request_uuid', None):
        req.request_uuid = f"iter-{req.iteration_number:04d}-{uuid.uuid4().hex[:6]}"

    backend = req.backend or LLM_BACKEND
    all_questions: List[Question] = []
    aggregated_metrics = {}

    try:
        if backend == 'local' and req.n_questions > LOCAL_GEN_BATCH_SIZE:
            remaining = req.n_questions
            batch_num = 0
            batch_total = (req.n_questions + LOCAL_GEN_BATCH_SIZE - 1) // LOCAL_GEN_BATCH_SIZE
            while remaining > 0:
                batch_num += 1
                n = min(LOCAL_GEN_BATCH_SIZE, remaining)
                logger.info(
                    "LLM batch %s/%s n=%s [Iteration #%s]",
                    batch_num,
                    batch_total,
                    n,
                    req.iteration_number,
                )
                batch_start = time.time()
                batch_req = req.model_copy(update={'n_questions': n})
                batch_questions = execute_generation_single(batch_req)
                batch_ms = int((time.time() - batch_start) * 1000)

                bm = getattr(batch_req, '_timing_metrics', {})
                if not aggregated_metrics:
                    aggregated_metrics = dict(bm)
                else:
                    for k in ['t_prompt_ms', 't_llm_ms', 't_validation_ms', 't_retry_ms', 't_gen_ms', 't_e2e_ms']:
                        aggregated_metrics[k] = round(aggregated_metrics.get(k, 0.0) + bm.get(k, 0.0), 2)
                    for k in ['validation_passed', 'validation_failed', 'retries_count']:
                        aggregated_metrics[k] = aggregated_metrics.get(k, 0) + bm.get(k, 0)

                log_generation(
                    request_uuid=f"{req.request_uuid}-b{batch_num}",
                    mode='sync_batch',
                    backend=backend,
                    model=getattr(req, 'model', None) or OLLAMA_MODEL_DEFAULT,
                    topic=req.topic,
                    level=req.level,
                    language=req.language,
                    n_questions_requested=n,
                    n_questions_generated=len(batch_questions),
                    duration_ms=batch_ms,
                    status='success',
                    batch_index=batch_num,
                    batch_total=batch_total,
                    has_lesson_context=bool(req.context),
                )
                all_questions.extend(batch_questions)
                remaining -= n

            req._timing_metrics = aggregated_metrics
        else:
            all_questions = execute_generation_single(req)
            aggregated_metrics = getattr(req, '_timing_metrics', {})

        # Log comprehensive evaluation iteration
        try:
            formatted_q = [
                {
                    'question': q.question,
                    'choices': [{'text': c.text, 'is_correct': c.is_correct} for c in q.choices],
                    'correct_index': q.correct_index,
                    'difficulty': q.difficulty,
                    'bloom_level': q.bloom_level,
                    'explanation': q.explanation,
                }
                for q in all_questions
            ]
            log_evaluation_iteration(
                iteration_number=req.iteration_number,
                request_uuid=req.request_uuid,
                topic=req.topic,
                difficulty=req.level,
                language=req.language,
                backend=backend,
                model=req.model or OLLAMA_MODEL_DEFAULT,
                category_name=getattr(req, 'category_name', '') or '',
                corpus_tokens=getattr(req, 'corpus_tokens', 0) or 0,
                chunk_count=aggregated_metrics.get('chunks_count', 0),
                top_k=getattr(req, 'top_k', 3) or 3,
                timing_metrics=aggregated_metrics,
                cache_hits=aggregated_metrics.get('cache_hits', 0),
                cache_misses=aggregated_metrics.get('cache_misses', 0),
                hit_ratio=aggregated_metrics.get('hit_ratio', 0.0),
                n_questions_requested=req.n_questions,
                n_questions_generated=len(all_questions),
                validation_passed=aggregated_metrics.get('validation_passed', len(all_questions)),
                validation_failed=aggregated_metrics.get('validation_failed', 0),
                retries_count=aggregated_metrics.get('retries_count', 0),
                questions=formatted_q,
                status='success' if all_questions else 'error',
            )
        except Exception as iter_err:
            logger.warning(f"Could not log evaluation iteration: {iter_err}")

        return all_questions

    except Exception as e:
        # Log failure evaluation iteration
        try:
            log_evaluation_iteration(
                iteration_number=req.iteration_number,
                request_uuid=req.request_uuid,
                topic=req.topic,
                difficulty=req.level,
                language=req.language,
                backend=backend,
                model=req.model or OLLAMA_MODEL_DEFAULT,
                category_name=getattr(req, 'category_name', '') or '',
                corpus_tokens=getattr(req, 'corpus_tokens', 0) or 0,
                chunk_count=0,
                top_k=getattr(req, 'top_k', 3) or 3,
                timing_metrics=aggregated_metrics or {},
                cache_hits=0,
                cache_misses=0,
                hit_ratio=0.0,
                n_questions_requested=req.n_questions,
                n_questions_generated=0,
                validation_passed=0,
                validation_failed=0,
                retries_count=0,
                questions=[],
                status='error',
                error_message=str(e),
            )
        except Exception:
            pass
        raise e


def execute_generation_single(req: QuestionRequest) -> List[Question]:
    backend = req.backend or LLM_BACKEND
    pipeline_mode = getattr(req, 'pipeline_mode', 'proposed') or 'proposed'
    top_k = getattr(req, 'top_k', 3) or 3

    # 1. RAG Vector Retrieval Step
    kb_metrics = {
        't_extract_ms': 0.0,
        't_chunk_ms': 0.0,
        't_hash_ms': 0.0,
        't_lookup_ms': 0.0,
        't_embed_ms': 0.0,
        't_index_ms': 0.0,
        't_kb_ms': 0.0,
        't_query_embed_ms': 0.0,
        't_retrieval_ms': 0.0,
        'cache_hits': 0,
        'cache_misses': 0,
        'hit_ratio': 0.0,
        'chunks_count': 0,
    }

    if req.context and req.topic and pipeline_mode != "no_rag":
        query = req.topic
        if req.bloom_level:
            query += " " + req.bloom_level
        if req.learning_outcomes:
            query += " " + req.learning_outcomes

        logger.info(f"Running RAG search with query: '{query}' [mode={pipeline_mode}, top_k={top_k}]")

        api_key = None
        if backend == 'openai':
            api_key = req.openai_api_key or OPENAI_API_KEY
        elif backend == 'gemini':
            api_key = req.gemini_api_key or GEMINI_API_KEY

        retrieved, kb_metrics = retrieve_relevant_context(
            query=query,
            document_text=req.context,
            backend=backend,
            api_key_override=api_key,
            local_model=req.model,
            top_k=top_k,
            max_context_chars=getattr(req, 'max_context_chars', None),
            enable_cache=bool(getattr(req, 'enable_incremental_cache', True)),
            pipeline_mode=pipeline_mode,
        )
        if retrieved:
            req.context = retrieved

    # 2. Generation phase with decomposed timing & code/AST validation
    from code_validator import validate_question

    t_prompt_ms = 0.0
    t_llm_ms = 0.0
    t_validation_ms = 0.0
    t_retry_ms = 0.0
    validation_passed = 0
    validation_failed = 0
    retries_count = 0

    max_retries = 3
    final_questions: List[Question] = []

    for attempt in range(max_retries):
        retries_count = attempt
        try:
            # T_prompt: Format template and outcomes
            t_p_start = time.perf_counter()
            _ = format_lesson_context(req.context)
            _ = format_learning_outcomes(req.learning_outcomes)
            t_prompt_ms += (time.perf_counter() - t_p_start) * 1000.0

            # T_LLM: LLM forward pass & token generation
            t_llm_start = time.perf_counter()
            questions = []
            q_type = getattr(req, 'question_type', 'code') or 'code'
            if backend == 'openai':
                effective_openai_key = req.openai_api_key or OPENAI_API_KEY
                if not effective_openai_key:
                    raise ValueError('OpenAI API key not configured')
                questions = generate_with_openai(
                    req.topic, req.level, req.n_questions,
                    req.language, req.bloom_level, req.context,
                    api_key_override=req.openai_api_key,
                    learning_outcomes=req.learning_outcomes,
                    question_type=q_type
                )
            elif backend == 'gemini':
                effective_gemini_key = req.gemini_api_key or GEMINI_API_KEY
                if not effective_gemini_key:
                    raise ValueError('Gemini API key not configured')
                questions = generate_with_gemini(
                    req.topic, req.level, req.n_questions,
                    req.language, req.bloom_level, req.context,
                    api_key_override=req.gemini_api_key,
                    learning_outcomes=req.learning_outcomes,
                    question_type=q_type
                )
            elif backend == 'local':
                questions = generate_with_local_llm(
                    req.topic, req.level, req.n_questions,
                    req.language, req.bloom_level, req.context,
                    model=req.model,
                    learning_outcomes=req.learning_outcomes,
                    question_type=q_type
                )
            else:
                raise ValueError(f'Unknown backend: {backend}')

            attempt_llm_ms = (time.perf_counter() - t_llm_start) * 1000.0
            t_llm_ms += attempt_llm_ms

            # 3. T_validation: Code AST validation and compiler checks
            t_val_start = time.perf_counter()
            valid_questions = []
            all_valid = True
            for q in questions:
                is_valid, err = validate_question(
                    question_text=q.question,
                    choices=[c.text for c in q.choices],
                    correct_index=q.correct_index,
                    topic=req.topic,
                    question_type=q_type
                )
                if is_valid:
                    valid_questions.append(q)
                    validation_passed += 1
                else:
                    validation_failed += 1
                    logger.warning(f"Generated question failed validation: {err}")
                    all_valid = False

            attempt_val_ms = (time.perf_counter() - t_val_start) * 1000.0
            t_validation_ms += attempt_val_ms

            if all_valid and len(valid_questions) == len(questions):
                final_questions = questions
                break

            if attempt == max_retries - 1:
                logger.warning(f"Max retries reached. Returning {len(valid_questions)} valid questions.")
                if not valid_questions:
                    raise ValueError("All generated questions failed compilation/syntax validation checks.")
                final_questions = valid_questions
                break

            # Add cost of failed attempt to retry latency
            t_retry_ms += (attempt_llm_ms + attempt_val_ms)
            logger.info(f"Retrying question generation (attempt {attempt + 2}/{max_retries}) due to compilation/validation errors...")

        except Exception as e:
            if attempt == max_retries - 1:
                raise e
            t_retry_ms += (time.perf_counter() - t_llm_start) * 1000.0 if 't_llm_start' in locals() else 0.0
            logger.warning(f"Error during attempt {attempt + 1}: {e}. Retrying...")

    # Calculate final decomposed metrics
    t_kb_ms = kb_metrics.get('t_kb_ms', 0.0)
    t_query_embed_ms = kb_metrics.get('t_query_embed_ms', 0.0)
    t_retrieval_ms = kb_metrics.get('t_retrieval_ms', 0.0)
    t_gen_ms = t_query_embed_ms + t_retrieval_ms + t_prompt_ms + t_llm_ms + t_validation_ms + t_retry_ms
    t_e2e_ms = t_kb_ms + t_gen_ms

    timing_metrics = {
        "pipeline_mode": pipeline_mode,
        "top_k": top_k,
        "corpus_tokens": getattr(req, 'corpus_tokens', 0) or 0,
        "change_ratio": getattr(req, 'change_ratio', 0.0) or 0.0,
        "t_extract_ms": round(kb_metrics.get('t_extract_ms', 0.0), 2),
        "t_chunk_ms": round(kb_metrics.get('t_chunk_ms', 0.0), 2),
        "t_hash_ms": round(kb_metrics.get('t_hash_ms', 0.0), 2),
        "t_lookup_ms": round(kb_metrics.get('t_lookup_ms', 0.0), 2),
        "t_embed_ms": round(kb_metrics.get('t_embed_ms', 0.0), 2),
        "t_index_ms": round(kb_metrics.get('t_index_ms', 0.0), 2),
        "t_kb_ms": round(t_kb_ms, 2),
        "t_query_embed_ms": round(t_query_embed_ms, 2),
        "t_retrieval_ms": round(t_retrieval_ms, 2),
        "t_prompt_ms": round(t_prompt_ms, 2),
        "t_llm_ms": round(t_llm_ms, 2),
        "t_validation_ms": round(t_validation_ms, 2),
        "t_retry_ms": round(t_retry_ms, 2),
        "t_gen_ms": round(t_gen_ms, 2),
        "t_e2e_ms": round(t_e2e_ms, 2),
        "cache_hits": kb_metrics.get('cache_hits', 0),
        "cache_misses": kb_metrics.get('cache_misses', 0),
        "hit_ratio": round(kb_metrics.get('hit_ratio', 0.0), 4),
        "validation_passed": validation_passed,
        "validation_failed": validation_failed,
        "retries_count": retries_count,
    }
    req._timing_metrics = timing_metrics

    # Log to research evaluation log
    try:
        log_inacon_metrics(
            request_uuid=getattr(req, 'request_uuid', '') or f"req-{int(time.time()*1000)}",
            pipeline_mode=pipeline_mode,
            backend=backend,
            model=req.model or OLLAMA_MODEL_DEFAULT,
            topic=req.topic,
            category_name=getattr(req, 'category_name', '') or '',
            corpus_tokens=getattr(req, 'corpus_tokens', 0) or 0,
            change_ratio=getattr(req, 'change_ratio', 0.0) or 0.0,
            top_k=top_k,
            t_extract_ms=timing_metrics['t_extract_ms'],
            t_chunk_ms=timing_metrics['t_chunk_ms'],
            t_hash_ms=timing_metrics['t_hash_ms'],
            t_lookup_ms=timing_metrics['t_lookup_ms'],
            t_embed_ms=timing_metrics['t_embed_ms'],
            t_index_ms=timing_metrics['t_index_ms'],
            t_kb_ms=timing_metrics['t_kb_ms'],
            t_query_embed_ms=timing_metrics['t_query_embed_ms'],
            t_retrieval_ms=timing_metrics['t_retrieval_ms'],
            t_prompt_ms=timing_metrics['t_prompt_ms'],
            t_llm_ms=timing_metrics['t_llm_ms'],
            t_validation_ms=timing_metrics['t_validation_ms'],
            t_retry_ms=timing_metrics['t_retry_ms'],
            t_gen_ms=timing_metrics['t_gen_ms'],
            t_e2e_ms=timing_metrics['t_e2e_ms'],
            cache_hits=timing_metrics['cache_hits'],
            cache_misses=timing_metrics['cache_misses'],
            hit_ratio=timing_metrics['hit_ratio'],
            n_questions_requested=req.n_questions,
            n_questions_generated=len(final_questions),
            validation_passed=validation_passed,
            validation_failed=validation_failed,
            retries_count=retries_count,
            status='success' if final_questions else 'error',
        )
    except Exception as log_err:
        logger.warning(f"Could not log pipeline metrics: {log_err}")

    return final_questions


def post_moodle_webhook(webhook_url: str, webhook_token: str, body: dict) -> None:
    """POST status/result to Moodle complete_generation_job.php."""
    headers = {
        'Content-Type': 'application/json',
        'X-Worker-Token': webhook_token,
    }
    logger.info(
        "Webhook → Moodle status=%s request_uuid=%s",
        body.get('status'),
        body.get('request_uuid'),
    )
    resp = requests.post(
        webhook_url,
        json=body,
        headers=headers,
        timeout=WEBHOOK_TIMEOUT,
    )
    if resp.status_code >= 400:
        raise RuntimeError(f'Moodle webhook HTTP {resp.status_code}: {resp.text[:300]}')
    data = resp.json()
    if data.get('success') is False:
        raise RuntimeError(data.get('error') or 'Moodle webhook reported failure')


def _async_generation_worker(payload: dict) -> None:
    """Background thread: generate questions and webhook Moodle."""
    request_uuid = payload['request_uuid']
    webhook_url = payload['webhook_url']
    webhook_token = payload['webhook_token']
    gen_start = time.time()

    try:
        req = QuestionRequest(**{k: v for k, v in payload.items() if k in QuestionRequest.model_fields})
        backend = req.backend or LLM_BACKEND

        post_moodle_webhook(webhook_url, webhook_token, {
            'request_uuid': request_uuid,
            'status': 'processing',
        })

        logger.info(
            "LLM async START request_uuid=%s backend=%s n=%s topic=%r",
            request_uuid,
            backend,
            req.n_questions,
            (req.topic or '')[:80],
        )

        questions = execute_generation(req)
        duration_ms = int((time.time() - gen_start) * 1000)

        log_generation(
            request_uuid=request_uuid,
            mode='async',
            backend=backend,
            model=payload.get('model') or OLLAMA_MODEL_DEFAULT,
            topic=req.topic,
            level=req.level,
            language=req.language,
            n_questions_requested=req.n_questions,
            n_questions_generated=len(questions),
            duration_ms=duration_ms,
            status='success',
            has_lesson_context=bool(req.context),
        )

        question_payload = []
        for q in questions:
            question_payload.append({
                'question': q.question,
                'choices': [{'text': c.text, 'is_correct': c.is_correct} for c in q.choices],
                'correct_index': q.correct_index,
                'difficulty': q.difficulty,
                'bloom_level': q.bloom_level,
                'explanation': q.explanation,
            })

        timing_metrics = getattr(req, '_timing_metrics', {})
        post_moodle_webhook(webhook_url, webhook_token, {
            'request_uuid': request_uuid,
            'status': 'success',
            'questions': question_payload,
            'generated_count': len(question_payload),
            'duration_ms': duration_ms,
            'iteration_number': getattr(req, 'iteration_number', 0),
            'iteration_id': f"iter_{getattr(req, 'iteration_number', 0):04d}",
            'timing_metrics': timing_metrics,
        })

        logger.info(
            "LLM async END request_uuid=%s count=%s duration=%.2fs",
            request_uuid,
            len(question_payload),
            duration_ms / 1000.0,
        )
    except Exception as e:
        duration_ms = int((time.time() - gen_start) * 1000)
        try:
            req = QuestionRequest(**{k: v for k, v in payload.items() if k in QuestionRequest.model_fields})
            log_generation(
                request_uuid=request_uuid,
                mode='async',
                backend=req.backend or LLM_BACKEND,
                model=payload.get('model') or OLLAMA_MODEL_DEFAULT,
                topic=req.topic,
                level=req.level,
                language=req.language,
                n_questions_requested=req.n_questions,
                n_questions_generated=0,
                duration_ms=duration_ms,
                status='error',
                error_message=str(e),
                has_lesson_context=bool(req.context),
            )
        except Exception:
            pass
        logger.error("LLM async FAIL request_uuid=%s: %s", request_uuid, e, exc_info=True)
        try:
            post_moodle_webhook(webhook_url, webhook_token, {
                'request_uuid': request_uuid,
                'status': 'error',
                'error_message': str(e),
                'duration_ms': duration_ms,
            })
        except Exception as webhook_err:
            logger.error(
                "Failed to send error webhook for %s: %s",
                request_uuid,
                webhook_err,
            )


@app.route('/generate/async', methods=['POST'])
def generate_questions_async():
    """Accept generation job and process in background; webhook Moodle when done."""
    try:
        data = request.json or {}
        async_req = AsyncGenerateRequest(**data)
        thread = threading.Thread(
            target=_async_generation_worker,
            args=(data,),
            daemon=True,
        )
        thread.start()
        logger.info(
            "LLM async ACCEPTED request_uuid=%s webhook=%s",
            async_req.request_uuid,
            async_req.webhook_url,
        )
        return jsonify({
            'status': 'accepted',
            'request_uuid': async_req.request_uuid,
            'message': 'Generation started; results will be sent via webhook',
        }), 202
    except Exception as e:
        logger.error(f"Error accepting async generation: {e}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/generate', methods=['POST'])
def generate_questions():
    """Generate MCQ questions (synchronous)."""
    try:
        data = request.json
        req = QuestionRequest(**data)

        backend = req.backend or LLM_BACKEND
        gen_start = time.time()
        logger.info(
            "LLM generate START backend=%s topic=%r n_questions=%s language=%s has_context=%s",
            backend,
            req.topic[:80] if req.topic else '',
            req.n_questions,
            req.language,
            bool(req.context),
        )

        questions = execute_generation(req)

        duration_s = time.time() - gen_start
        duration_ms = int(duration_s * 1000)
        log_generation(
            request_uuid=data.get('request_uuid', '') or f"sync-{int(gen_start)}",
            mode='sync',
            backend=backend,
            model=data.get('model') or OLLAMA_MODEL_DEFAULT,
            topic=req.topic,
            level=req.level,
            language=req.language,
            n_questions_requested=req.n_questions,
            n_questions_generated=len(questions),
            duration_ms=duration_ms,
            status='success',
            has_lesson_context=bool(req.context),
        )
        logger.info(
            "LLM generate END backend=%s count=%s duration=%.2fs",
            backend,
            len(questions),
            duration_s,
        )

        timing_metrics = getattr(req, '_timing_metrics', {})
        response = QuestionResponse(
            questions=questions,
            metadata={
                'topic': req.topic,
                'language': req.language,
                'count': len(questions),
                'backend': backend,
                'iteration_number': getattr(req, 'iteration_number', 0),
                'iteration_id': f"iter_{getattr(req, 'iteration_number', 0):04d}",
                'timing_metrics': timing_metrics,
            }
        )

        return jsonify(response.model_dump()), 200

    except Exception as e:
        logger.error(f"Error generating questions: {str(e)}", exc_info=True)
        try:
            data = request.json or {}
            req = QuestionRequest(**data)
            log_generation(
                request_uuid=data.get('request_uuid', '') or 'sync-error',
                mode='sync',
                backend=req.backend or LLM_BACKEND,
                model=data.get('model') or OLLAMA_MODEL_DEFAULT,
                topic=req.topic,
                level=req.level,
                language=req.language,
                n_questions_requested=req.n_questions,
                n_questions_generated=0,
                duration_ms=0,
                status='error',
                error_message=str(e),
                has_lesson_context=bool(req.context),
            )
        except Exception:
            pass
        return jsonify({'error': str(e)}), 500


@app.route('/validate', methods=['POST'])
def validate_question():
    """Validate question quality"""
    # Placeholder for question validation
    data = request.json
    return jsonify({
        'valid': True,
        'score': 0.85,
        'feedback': 'Question quality is good'
    }), 200


def extract_text_from_file_bytes(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from PDF, PPTX, DOCX, or text files."""
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == '.pdf':
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(f"--- Page {i+1} ---\n{text.strip()}")
            return "\n\n".join(pages)
        except ImportError:
            raise RuntimeError("pypdf is not installed. Please run 'pip install pypdf'.")
            
    elif ext in ('.pptx', '.ppt'):
        try:
            import pptx
            prs = pptx.Presentation(io.BytesIO(file_bytes))
            slides = []
            for i, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    slides.append(f"--- Slide {i+1} ---\n" + "\n".join(slide_text))
            return "\n\n".join(slides)
        except ImportError:
            raise RuntimeError("python-pptx is not installed. Please run 'pip install python-pptx'.")
            
    elif ext in ('.docx', '.doc'):
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)
        except ImportError:
            raise RuntimeError("python-docx is not installed. Please run 'pip install python-docx'.")
            
    elif ext in ('.png', '.jpg', '.jpeg', '.webp', '.bmp', '.gif'):
        try:
            import base64
            img_b64 = base64.b64encode(file_bytes).decode('utf-8')
            resp = requests.post(f"{LOCAL_LLM_URL}/api/generate", json={
                "model": "qwen2.5vl:3b",
                "prompt": "Transcribe all text, code, formulas, diagrams, labels, and key educational concepts visible in this image clearly and concisely.",
                "images": [img_b64],
                "stream": False
            }, timeout=60)
            if resp.status_code == 200:
                transcription = resp.json().get('response', '').strip()
                if transcription:
                    return f"=== Image: {filename} ===\n{transcription}"
            return f"[Image: {filename}]"
        except Exception as e:
            logger.warning(f"Image vision transcription error for {filename}: {e}")
            return f"[Image: {filename}]"

    else:
        # Default plain text decode for .txt, .md, .py, .c, .cpp, .java, .json, .csv, .html
        return file_bytes.decode('utf-8', errors='replace')


@app.route('/extract_file', methods=['POST'])
def extract_file_endpoint():
    """Extract text from uploaded PDF, PPTX, DOCX, or plain text course files."""
    filename = "document"
    try:
        file_bytes = b""
        
        # 1. Check multipart/form-data upload
        if 'file' in request.files:
            uploaded_file = request.files['file']
            filename = uploaded_file.filename or "uploaded_file"
            file_bytes = uploaded_file.read()
        # 2. Check JSON payload with base64 or file_path
        elif request.is_json:
            data = request.json or {}
            filename = data.get('filename', 'document.txt')
            if 'file_content_base64' in data:
                file_bytes = base64.b64decode(data['file_content_base64'])
            elif 'file_path' in data:
                file_path = data['file_path']
                if os.path.isfile(file_path):
                    filename = os.path.basename(file_path)
                    with open(file_path, 'rb') as f:
                        file_bytes = f.read()
                else:
                    return jsonify({'success': False, 'error': f'File not found: {file_path}'}), 404
        
        if not file_bytes:
            return jsonify({'success': False, 'error': 'No file received. Upload multipart file or provide JSON file_content_base64'}), 400
            
        extracted_text = extract_text_from_file_bytes(file_bytes, filename)
        chars = len(extracted_text)
        approx_tokens = int(chars / 4)
        
        return jsonify({
            'success': True,
            'filename': filename,
            'text': extracted_text,
            'characters': chars,
            'approx_tokens': approx_tokens
        }), 200
    except Exception as e:
        logger.error(f"Error extracting text from file {filename}: {str(e)}")
        return jsonify({'success': False, 'error': str(e)}), 500


def check_content_cache_status(content: str, backend: str = LLM_BACKEND, model: str = OLLAMA_MODEL_DEFAULT) -> dict:
    """Helper to check SHA-256 chunk caching status for a given text."""
    content = (content or '').strip()
    if not content:
        return {
            'has_content': False,
            'status': 'empty',
            'total_chunks': 0,
            'cached_chunks': 0,
            'modified_chunks': 0,
            'is_synced': True,
            'content_hash': '',
            'short_hash': '',
            'message': 'No content provided.'
        }

    content_hash = hashlib.sha256(content.encode('utf-8')).hexdigest()
    chunks = chunk_text(content)

    if backend == 'local':
        model_name = os.getenv('OLLAMA_EMBED_MODEL', 'nomic-embed-text') or model
    elif backend == 'openai':
        model_name = "text-embedding-3-small"
    elif backend == 'gemini':
        model_name = "models/text-embedding-004"
    else:
        model_name = "unknown"

    cached_count = 0
    with embedding_cache_lock:
        for c in chunks:
            hash_key = hashlib.sha256(f"{model_name}:{c}".encode('utf-8')).hexdigest()
            if hash_key in embedding_cache:
                cached_count += 1

    total_chunks = len(chunks)
    modified_chunks = total_chunks - cached_count
    is_synced = (modified_chunks == 0 and total_chunks > 0)

    if is_synced:
        status = 'synced'
        msg = f"Embedding Cache Synced: All {total_chunks} chunks cached (SHA-256: {content_hash[:8]}). Sub-millisecond retrieval ready."
    elif cached_count > 0:
        status = 'modified'
        msg = f"Notice: Lesson has changed since last cache index. {modified_chunks} of {total_chunks} chunks require re-indexing."
    else:
        status = 'uncached'
        msg = f"Notice: Lesson has not been indexed in embedding cache yet ({total_chunks} chunks ready to index)."

    return {
        'has_content': True,
        'status': status,
        'content_hash': content_hash,
        'short_hash': content_hash[:8],
        'total_chunks': total_chunks,
        'cached_chunks': cached_count,
        'modified_chunks': modified_chunks,
        'is_synced': is_synced,
        'message': msg
    }


@app.route('/cache/status', methods=['POST'])
def cache_status():
    """Check SHA-256 chunk caching status for a given lesson text."""
    try:
        data = request.json or {}
        content = data.get('content', '').strip()
        backend = data.get('backend', LLM_BACKEND)
        model = data.get('model', OLLAMA_MODEL_DEFAULT)

        res = check_content_cache_status(content, backend, model)
        return jsonify(res)
    except Exception as e:
        logger.error(f"Error checking cache status: {e}", exc_info=True)
        return jsonify({'error': str(e)}), 500


@app.route('/cache/batch_status', methods=['POST'])
def cache_batch_status():
    """Check SHA-256 chunk caching status for multiple lesson texts at once."""
    try:
        data = request.json or {}
        items = data.get('items', {})
        backend = data.get('backend', LLM_BACKEND)
        model = data.get('model', OLLAMA_MODEL_DEFAULT)

        results = {}
        all_synced = True
        synced_count = 0
        modified_count = 0
        uncached_count = 0
        empty_count = 0

        for key, text in items.items():
            st = check_content_cache_status(text, backend, model)
            results[key] = st
            if st['status'] == 'synced':
                synced_count += 1
            elif st['status'] == 'modified':
                modified_count += 1
                all_synced = False
            elif st['status'] == 'uncached':
                uncached_count += 1
                all_synced = False
            elif st['status'] == 'empty':
                empty_count += 1

        return jsonify({
            'success': True,
            'sources': results,
            'all_synced': (modified_count == 0 and uncached_count == 0),
            'total_sources': len(items),
            'synced_sources': synced_count,
            'modified_sources': modified_count,
            'uncached_sources': uncached_count,
            'empty_sources': empty_count
        })
    except Exception as e:
        logger.error(f"Error checking batch cache status: {e}", exc_info=True)
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/cache/details', methods=['POST'])
def cache_details():
    """Return in-depth chunking, embedding vector samples, and cache metrics for a source."""
    try:
        data = request.json or {}
        content = (data.get('content') or '').strip()
        source_id = data.get('source_id', '')
        backend = data.get('backend', LLM_BACKEND)
        model = data.get('model', OLLAMA_MODEL_DEFAULT)

        if backend == 'local':
            embed_model_name = os.getenv('OLLAMA_EMBED_MODEL', 'nomic-embed-text') or 'nomic-embed-text'
            default_dim = 768
        elif backend == 'openai':
            embed_model_name = "text-embedding-3-small"
            default_dim = 1536
        elif backend == 'gemini':
            embed_model_name = "models/text-embedding-004"
            default_dim = 768
        else:
            embed_model_name = "nomic-embed-text"
            default_dim = 768

        if not content:
            return jsonify({
                'success': True,
                'source_id': source_id,
                'has_content': False,
                'status': 'empty',
                'content_hash': '',
                'short_hash': '',
                'total_chars': 0,
                'estimated_tokens': 0,
                'total_chunks': 0,
                'cached_chunks': 0,
                'modified_chunks': 0,
                'is_synced': True,
                'embedding_model': embed_model_name,
                'vector_dim': default_dim,
                'distance_metric': 'Cosine Similarity',
                'chunks': [],
                'raw_content': '',
                'message': 'No content extractable for this source.'
            })

        content_hash = hashlib.sha256(content.encode('utf-8')).hexdigest()
        chunks = chunk_text(content)

        chunks_detail = []
        cached_count = 0

        with embedding_cache_lock:
            for idx, c in enumerate(chunks):
                hash_key = hashlib.sha256(f"{embed_model_name}:{c}".encode('utf-8')).hexdigest()
                is_cached = hash_key in embedding_cache
                vec = embedding_cache.get(hash_key)

                if is_cached and vec:
                    cached_count += 1
                    vec_len = len(vec)
                    vec_sample = [round(float(x), 4) for x in vec[:10]]
                    vec_norm = round(sum(float(x)**2 for x in vec)**0.5, 4)
                    vec_full = [round(float(x), 5) for x in vec]
                else:
                    vec_len = default_dim
                    vec_sample = []
                    vec_norm = 0.0
                    vec_full = []

                chunks_detail.append({
                    'index': idx + 1,
                    'sha256': hash_key,
                    'short_sha256': hash_key[:12],
                    'char_length': len(c),
                    'estimated_tokens': max(1, round(len(c) / 4)),
                    'is_cached': is_cached,
                    'text': c,
                    'vector_dim': vec_len,
                    'vector_sample': vec_sample,
                    'vector_norm': vec_norm,
                    'vector_full': vec_full
                })

        total_chunks = len(chunks)
        modified_chunks = total_chunks - cached_count
        is_synced = (modified_chunks == 0 and total_chunks > 0)
        status = 'synced' if is_synced else ('modified' if cached_count > 0 else 'uncached')

        return jsonify({
            'success': True,
            'source_id': source_id,
            'has_content': True,
            'status': status,
            'content_hash': content_hash,
            'short_hash': content_hash[:8],
            'total_chars': len(content),
            'estimated_tokens': max(1, round(len(content) / 4)),
            'total_chunks': total_chunks,
            'cached_chunks': cached_count,
            'modified_chunks': modified_chunks,
            'is_synced': is_synced,
            'embedding_model': embed_model_name,
            'vector_dim': default_dim,
            'distance_metric': 'Cosine Similarity',
            'chunks': chunks_detail,
            'raw_content': content,
            'message': f"Analysis complete: {cached_count}/{total_chunks} chunks cached."
        })
    except Exception as e:
        logger.error(f"Error fetching cache details: {e}", exc_info=True)
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/cache/reindex', methods=['POST'])
def cache_reindex():
    """Force recompute and persist SHA-256 chunk embeddings into cache (single or batch)."""
    try:
        data = request.json or {}
        content = data.get('content', '').strip()
        items = data.get('items', {})
        backend = data.get('backend', LLM_BACKEND)
        model = data.get('model', OLLAMA_MODEL_DEFAULT)
        api_key_override = data.get('api_key')

        if not content and not items:
            return jsonify({'error': 'No content provided to index.'}), 400

        t_start = time.perf_counter()

        if backend == 'local':
            model_name = os.getenv('OLLAMA_EMBED_MODEL', 'nomic-embed-text') or model
        elif backend == 'openai':
            model_name = "text-embedding-3-small"
        elif backend == 'gemini':
            model_name = "models/text-embedding-004"
        else:
            model_name = "unknown"

        def compute_single_embedding(text: str) -> list:
            try:
                if backend == 'openai':
                    from openai import OpenAI
                    client = OpenAI(api_key=api_key_override or OPENAI_API_KEY)
                    resp = client.embeddings.create(input=[text], model=model_name)
                    return resp.data[0].embedding
                elif backend == 'gemini':
                    import google.generativeai as genai
                    genai.configure(api_key=api_key_override or GEMINI_API_KEY)
                    resp = genai.embed_content(model=model_name, content=text)
                    return resp['embedding']
                elif backend == 'local':
                    resp = requests.post(f"{LOCAL_LLM_URL}/api/embeddings", json={"model": model_name, "prompt": text}, timeout=60)
                    if resp.status_code == 200:
                        return resp.json().get('embedding', [])
                    return []
                return []
            except Exception as e:
                logger.error(f"Embedding error: {e}")
                return []

        # If items dict provided, process each source item in batch
        if items:
            total_chunks_all = 0
            newly_indexed_all = 0
            new_entries = {}

            for sid, text_content in items.items():
                txt = (text_content or '').strip()
                if not txt:
                    continue

                chunks = chunk_text(txt)
                total_chunks_all += len(chunks)
                for c in chunks:
                    hash_key = hashlib.sha256(f"{model_name}:{c}".encode('utf-8')).hexdigest()
                    with embedding_cache_lock:
                        if hash_key in embedding_cache or hash_key in new_entries:
                            continue
                    vec = compute_single_embedding(c)
                    if vec:
                        new_entries[hash_key] = vec
                        newly_indexed_all += 1

            if new_entries:
                with embedding_cache_lock:
                    embedding_cache.update(new_entries)
                    save_embedding_cache()

            source_results = {}
            for sid, text_content in items.items():
                source_results[sid] = check_content_cache_status(text_content, backend, model)

            duration_ms = round((time.perf_counter() - t_start) * 1000.0, 2)
            return jsonify({
                'success': True,
                'sources': source_results,
                'total_chunks': total_chunks_all,
                'newly_indexed': newly_indexed_all,
                'duration_ms': duration_ms,
                'message': f"Batch indexed {len(items)} sources ({newly_indexed_all} new chunks, {duration_ms} ms)."
            })

        # Single content reindexing
        content_hash = hashlib.sha256(content.encode('utf-8')).hexdigest()
        chunks = chunk_text(content)

        new_entries = {}
        newly_indexed = 0

        for c in chunks:
            hash_key = hashlib.sha256(f"{model_name}:{c}".encode('utf-8')).hexdigest()
            with embedding_cache_lock:
                if hash_key in embedding_cache:
                    continue
            vec = compute_single_embedding(c)
            if vec:
                new_entries[hash_key] = vec
                newly_indexed += 1

        if new_entries:
            with embedding_cache_lock:
                embedding_cache.update(new_entries)
                save_embedding_cache()

        duration_ms = round((time.perf_counter() - t_start) * 1000.0, 2)
        total_chunks = len(chunks)

        with embedding_cache_lock:
            cached_count = sum(1 for c in chunks if hashlib.sha256(f"{model_name}:{c}".encode('utf-8')).hexdigest() in embedding_cache)

        modified_chunks = total_chunks - cached_count
        is_synced = (modified_chunks == 0 and total_chunks > 0)
        status = 'synced' if is_synced else ('modified' if cached_count > 0 else 'uncached')

        return jsonify({
            'success': is_synced,
            'status': status,
            'is_synced': is_synced,
            'content_hash': content_hash,
            'short_hash': content_hash[:8],
            'total_chunks': total_chunks,
            'cached_chunks': cached_count,
            'modified_chunks': modified_chunks,
            'newly_indexed': newly_indexed,
            'duration_ms': duration_ms,
            'message': f"All {total_chunks} chunks indexed into SHA-256 cache ({newly_indexed} new, {duration_ms} ms, Hash: {content_hash[:8]})." if is_synced else f"Indexed {cached_count}/{total_chunks} chunks ({newly_indexed} new, {duration_ms} ms)."
        })
    except Exception as e:
        logger.error(f"Error reindexing cache: {e}", exc_info=True)
        return jsonify({'error': str(e)}), 500


# -------------------------------------------------------------------------
# Evaluation & Telemetry Dashboard Endpoints
# -------------------------------------------------------------------------

@app.route('/dashboard')
def evaluation_dashboard():
    """Render interactive HTML evaluation dashboard."""
    return render_template('dashboard.html')


@app.route('/evaluation')
@app.route('/evaluate')
def research_evaluation_dashboard():
    """Render the comprehensive research evaluation dashboard (E1-E4)."""
    eval_html_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'evaluate', 'dashboard.html')
    if os.path.exists(eval_html_path):
        with open(eval_html_path, 'r', encoding='utf-8') as f:
            return f.read(), 200, {'Content-Type': 'text/html; charset=utf-8'}
    return "Evaluation dashboard not generated yet. Run scripts/build_evaluation_dashboard.py", 404


@app.route('/api/dashboard/stats', methods=['GET'])
def api_dashboard_stats():
    """Global aggregate stats across all recorded iterations."""
    return jsonify(get_evaluation_dashboard_stats()), 200


@app.route('/api/dashboard/iterations', methods=['GET'])
def api_dashboard_iterations():
    """List summary for all iterations (for dropdown selector and comparison table)."""
    limit = int(request.args.get('limit', 500))
    return jsonify({'iterations': get_all_iterations_summary(limit=limit)}), 200


@app.route('/api/dashboard/iteration/<int:iter_num>', methods=['GET'])
def api_dashboard_iteration_detail(iter_num: int):
    """Detailed logs for a specific single iteration."""
    details = get_iteration_details(iter_num)
    if not details:
        return jsonify({'error': f'Iteration #{iter_num} not found'}), 404
    return jsonify(details), 200


@app.route('/api/dashboard/reset', methods=['POST'])
def api_dashboard_reset():
    """Reset the iteration counter to a specified number."""
    data = request.json or {}
    start = int(data.get('start', 1))
    reset_iteration_number(start)
    return jsonify({'success': True, 'current_iteration': start - 1}), 200


@app.route('/api/dashboard/current_iteration', methods=['GET'])
def api_dashboard_current_iteration():
    """Return the current active iteration counter."""
    return jsonify({'current_iteration': get_current_iteration_number()}), 200



if __name__ == '__main__':
    _preload_ollama_models()
    port = int(os.getenv('FLASK_PORT', 5001))
    app.run(host='0.0.0.0', port=port, debug=os.getenv('NODE_ENV') == 'development')

