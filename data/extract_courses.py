#!/usr/bin/env python3
"""
Extract text and structured lecture units from authentic course materials in data/courses.
Extracts from PDF, DOCX, and PPTX files.
"""

import os
import sys
import json
import pypdf
import docx
import pptx

def extract_pdf(path: str) -> str:
    text = []
    reader = pypdf.PdfReader(path)
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            text.append(f"--- Page {i+1} ---\n{page_text.strip()}")
    return "\n\n".join(text)

def extract_docx(path: str) -> str:
    doc = docx.Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)

def extract_pptx(path: str) -> str:
    prs = pptx.Presentation(path)
    slides_text = []
    for i, slide in enumerate(prs.slides):
        slide_parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())
        if slide_parts:
            slides_text.append(f"--- Slide {i+1} ---\n" + "\n".join(slide_parts))
    return "\n\n".join(slides_text)

def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    courses_dir = os.path.join(base_dir, "courses")
    out_dir = os.path.join(base_dir, "extracted")
    os.makedirs(out_dir, exist_ok=True)
    # Clean stale files
    for old_f in os.listdir(out_dir):
        if old_f.endswith(".txt") or old_f.endswith(".json"):
            try:
                os.remove(os.path.join(out_dir, old_f))
            except Exception:
                pass

    manifest = []
    all_corpus_parts = []

    print(f"=== Extracting Course Materials from {courses_dir} ===")

    for fname in sorted(os.listdir(courses_dir)):
        fpath = os.path.join(courses_dir, fname)
        if not os.path.isfile(fpath) or fname.startswith("."):
            continue

        raw_text = ""
        ext = os.path.splitext(fname)[1].lower()
        if ext == ".pdf":
            raw_text = extract_pdf(fpath)
        elif ext == ".docx":
            raw_text = extract_docx(fpath)
        elif ext == ".pptx":
            raw_text = extract_pptx(fpath)
        else:
            print(f"[SKIP] Unsupported file format: {fname}")
            continue

        safe_name = os.path.splitext(fname)[0].replace(" ", "_").replace(",", "").lower()
        txt_out = os.path.join(out_dir, f"{safe_name}.txt")
        with open(txt_out, "w", encoding="utf-8") as f:
            f.write(raw_text)

        chars = len(raw_text)
        words = len(raw_text.split())
        approx_tokens = int(chars / 4)

        manifest.append({
            "source_file": fname,
            "extracted_txt": f"{safe_name}.txt",
            "characters": chars,
            "words": words,
            "approx_tokens": approx_tokens
        })

        header = f"\n\n{'='*70}\nCOURSE MODULE: {os.path.splitext(fname)[0]}\n{'='*70}\n\n"
        all_corpus_parts.append(header + raw_text)

        print(f"[EXTRACTED] {fname} -> {chars:,} chars (~{approx_tokens:,} tokens)")

    full_corpus = "".join(all_corpus_parts).strip()
    full_corpus_path = os.path.join(out_dir, "full_course_corpus.txt")
    with open(full_corpus_path, "w", encoding="utf-8") as f:
        f.write(full_corpus)

    manifest_path = os.path.join(out_dir, "manifest.json")
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)

    total_chars = len(full_corpus)
    total_tokens = int(total_chars / 4)
    print("\n[SUCCESS] Extraction Complete!")
    print(f"Total Modules Extracted: {len(manifest)}")
    print(f"Total Corpus Characters: {total_chars:,}")
    print(f"Estimated Total Tokens:  {total_tokens:,}")
    print(f"Full Corpus Saved to:    {full_corpus_path}")

if __name__ == "__main__":
    main()
