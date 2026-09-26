import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_formatted_text(paragraph, text, base_font_size=11, is_italic=False, is_bold=False):
    # Regex to tokenize bold, italic, code, and inline math
    # Tokens: `code`, **bold**, *italic*, $math$
    pattern = r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\$[^\$]+?\$)'
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = 'Times New Roman'
        run.font.size = Pt(base_font_size)
        run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
        
        if part.startswith('**') and part.endswith('**'):
            run.text = part[2:-2]
            run.bold = True
            run.italic = is_italic
        elif part.startswith('*') and part.endswith('*'):
            run.text = part[1:-1]
            run.italic = True
            run.bold = is_bold
        elif part.startswith('`') and part.endswith('`'):
            run.text = part[1:-1]
            run.font.name = 'Courier New'
            run.font.size = Pt(base_font_size - 0.5)
            run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
        elif part.startswith('$') and part.endswith('$'):
            # Math formatting
            run.text = part[1:-1]
            run.font.name = 'Cambria Math'
            run.italic = True
        else:
            run.text = part
            run.bold = is_bold
            run.italic = is_italic

def convert_md_to_docx(md_path, docx_path):
    doc = Document()
    
    # Page Margins: 1 inch (72 pt) all sides
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    i = 0
    n = len(lines)
    
    while i < n:
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()
        
        # Blank line
        if not stripped:
            i += 1
            continue
            
        # Title (# Title)
        if stripped.startswith('# '):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(14)
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
            i += 1
            continue
            
        # Author / Affiliation blocks
        if stripped.startswith('<div class="author-block">'):
            author_text = re.sub(r'<[^>]+>', '', stripped).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(author_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True
            i += 1
            continue
            
        if stripped.startswith('<div class="affil-block">'):
            affil_text = re.sub(r'<br\s*/?>', '\n', stripped)
            affil_text = re.sub(r'<[^>]+>', '', affil_text).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(14)
            for subline in affil_text.split('\n'):
                run = p.add_run(subline + '\n')
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                run.italic = True
                run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            i += 1
            continue
            
        if stripped == '---':
            i += 1
            continue
            
        # Heading 1 (## Heading)
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
            i += 1
            continue
            
        # Heading 2 (### Heading)
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
            i += 1
            continue

        # Heading 3 (#### Heading)
        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.bold = True
            run.italic = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # Image ![caption](path)
        if stripped.startswith('![') and '](' in stripped:
            img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
            if img_match:
                caption = img_match.group(1)
                img_path = img_match.group(2)
                # Map SVG to generated PNG
                if 'pipeline_architecture' in img_path:
                    png_path = '/Users/engtitya/Desktop/kwiz/papers/figures/pipeline_architecture.png'
                elif 'cache_decision_flow' in img_path:
                    png_path = '/Users/engtitya/Desktop/kwiz/papers/figures/cache_decision_flow.png'
                else:
                    png_path = os.path.join('/Users/engtitya/Desktop/kwiz/papers', img_path)
                
                if os.path.exists(png_path):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(8)
                    p_img.paragraph_format.space_after = Pt(2)
                    run_img = p_img.add_run()
                    run_img.add_picture(png_path, width=Inches(6.2))
                    
                i += 1
                continue

        # Figure caption (*Figure X: ...*)
        if stripped.startswith('*Figure ') and stripped.endswith('*'):
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
            run_cap = p_cap.add_run(stripped[1:-1])
            run_cap.font.name = 'Times New Roman'
            run_cap.font.size = Pt(9.5)
            run_cap.italic = True
            run_cap.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            i += 1
            continue

        # Blockquote (> text)
        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_formatted_text(p, stripped[2:], base_font_size=10.5, is_italic=True)
            i += 1
            continue

        # Display Math ($$...$$)
        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4:
            p_math = doc.add_paragraph()
            p_math.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_math.paragraph_format.space_before = Pt(6)
            p_math.paragraph_format.space_after = Pt(6)
            run = p_math.add_run(stripped[2:-2].strip())
            run.font.name = 'Cambria Math'
            run.font.size = Pt(11)
            run.italic = True
            i += 1
            continue

        # Tables (| col1 | col2 |)
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
                
            if len(table_lines) >= 2:
                # First line: headers
                headers = [c.strip() for c in table_lines[0][1:-1].split('|')]
                # Second line: alignment / separator
                # Data lines:
                data_rows = []
                for tl in table_lines[2:]:
                    cols = [c.strip() for c in tl[1:-1].split('|')]
                    data_rows.append(cols)
                    
                table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True
                
                # Format header row
                hdr_cells = table.rows[0].cells
                for col_idx, h_text in enumerate(headers):
                    hdr_cells[col_idx].text = ""
                    p = hdr_cells[col_idx].paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                    add_formatted_text(p, h_text, base_font_size=9.5, is_bold=True)
                    set_cell_background(hdr_cells[col_idx], "F1F5F9")
                    set_cell_margins(hdr_cells[col_idx], top=80, bottom=80, left=120, right=120)
                    
                # Format data rows
                for r_idx, row_data in enumerate(data_rows):
                    row_cells = table.rows[r_idx + 1].cells
                    bg_color = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
                    for col_idx, cell_value in enumerate(row_data):
                        if col_idx < len(row_cells):
                            row_cells[col_idx].text = ""
                            p = row_cells[col_idx].paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            # Align center for numbers, left for labels
                            is_center = col_idx > 0 or cell_value.isdigit()
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_center else WD_ALIGN_PARAGRAPH.LEFT
                            add_formatted_text(p, cell_value, base_font_size=9)
                            set_cell_background(row_cells[col_idx], bg_color)
                            set_cell_margins(row_cells[col_idx], top=60, bottom=60, left=100, right=100)
                            
                # Add spacing after table
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_before = Pt(0)
                p_after.paragraph_format.space_after = Pt(6)
            continue

        # Bullet list (- or * or •)
        if re.match(r'^[-*•]\s+', stripped):
            item_text = re.sub(r'^[-*•]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, item_text, base_font_size=11)
            i += 1
            continue

        # Numbered list (1. 2. etc)
        if re.match(r'^\d+\.\s+', stripped):
            item_text = re.sub(r'^\d+\.\s+', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, item_text, base_font_size=11)
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        # Check if it is the References section
        if stripped.startswith('[') and ']' in stripped and re.match(r'^\[\d+\]', stripped):
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            add_formatted_text(p, stripped, base_font_size=9.5)
        else:
            add_formatted_text(p, stripped, base_font_size=11)
            
        i += 1

    doc.save(docx_path)
    print(f"Generated {docx_path} ({os.path.getsize(docx_path)} bytes)")

if __name__ == '__main__':
    md_file = '/Users/engtitya/Desktop/kwiz/papers/paper.md'
    docx_file = '/Users/engtitya/Desktop/kwiz/papers/paper.docx'
    convert_md_to_docx(md_file, docx_file)
