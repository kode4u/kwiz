import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

import sys
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from convert_to_word import add_math_runs, add_formatted_text, set_cell_background, set_cell_margins

def create_eait_title_page(output_path):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_title.add_run("Toward Efficient Course-Grounded Programming MCQ Generation: An End-to-End Self-Hosted RAG Pipeline for Moodle")
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
    
    # Subtitle / Note
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Title Page for Double-Blind Review\nTarget Journal: Education and Information Technologies (EAIT)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(10)
    r_sub.italic = True
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    # Author Information
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(12)
    p_h1.paragraph_format.space_after = Pt(4)
    r = p_h1.add_run("Author Information")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    
    p_auth = doc.add_paragraph()
    p_auth.paragraph_format.space_before = Pt(2)
    p_auth.paragraph_format.space_after = Pt(4)
    r = p_auth.add_run("Author: ")
    r.bold = True
    r.font.name = "Times New Roman"
    r2 = p_auth.add_run("ENG Titya\n")
    r2.font.name = "Times New Roman"
    
    r_aff = p_auth.add_run("Affiliation: ")
    r_aff.bold = True
    r_aff.font.name = "Times New Roman"
    r_aff2 = p_auth.add_run("Faculty of Science and Technology, National University of Battambang, Battambang 020101, Cambodia\n")
    r_aff2.font.name = "Times New Roman"
    
    r_cor = p_auth.add_run("Corresponding Author: ")
    r_cor.bold = True
    r_cor.font.name = "Times New Roman"
    r_cor2 = p_auth.add_run("ENG Titya (Email: eng.titya@nubb.edu.kh)\n")
    r_cor2.font.name = "Times New Roman"

    # Declarations
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(16)
    p_h2.paragraph_format.space_after = Pt(4)
    r = p_h2.add_run("Statements and Declarations")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    
    decls = [
        ("Funding", "This research was supported by INACON research funding."),
        ("Competing Interests", "The author declares that they have no competing financial or non-financial interests that are directly or indirectly related to the work submitted for publication."),
        ("Ethics Approval", "Not applicable. This investigation evaluates an automated computational software pipeline, local machine learning models, and synthetic assessment question generation. No human subjects, patients, or student academic records were involved, and no personal identifiable data were collected or processed."),
        ("Informed Consent to Participate", "Not applicable."),
        ("Consent for Publication", "Not applicable."),
        ("Data Availability", "All course benchmark materials, configuration templates, rating rubrics, and empirical telemetry logs generated during the study are included in the companion open-source repository."),
        ("Code Availability", "The complete self-hosted RAG pipeline, Moodle Question Bank integration plugin, Docker Compose multi-container orchestrations, and automated evaluation suites are publicly available at https://github.com/kode4u/kwiz.git."),
        ("Author Contributions", "ENG Titya conceived and designed the study, implemented the RAG pipeline architecture and native Moodle plugin integration, performed the experimental benchmarks and statistical evaluations, analyzed the empirical results, and authored the complete manuscript."),
        ("Generative AI Disclosure", "In accordance with Springer Nature editorial policy, the author confirms that large language models (local Qwen2.5-Coder-7B and frontier LLM judges) served as the operational subject of the research pipeline and that automated LLM evaluation was calibrated against deterministic AST verification. Generative tools were only utilized for minor English language proofreading; the author takes full personal responsibility for the accuracy and originality of all content.")
    ]
    
    for h, text in decls:
        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_before = Pt(3)
        p_d.paragraph_format.space_after = Pt(4)
        p_d.paragraph_format.line_spacing = 1.15
        r_bold = p_d.add_run(f"• {h}: ")
        r_bold.font.name = "Times New Roman"
        r_bold.font.size = Pt(11)
        r_bold.bold = True
        r_txt = p_d.add_run(text)
        r_txt.font.name = "Times New Roman"
        r_txt.font.size = Pt(11)

    # Acknowledgments
    p_h3 = doc.add_paragraph()
    p_h3.paragraph_format.space_before = Pt(16)
    p_h3.paragraph_format.space_after = Pt(4)
    r = p_h3.add_run("Acknowledgments")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    
    p_ack = doc.add_paragraph()
    p_ack.paragraph_format.space_before = Pt(2)
    p_ack.paragraph_format.space_after = Pt(6)
    p_ack.paragraph_format.line_spacing = 1.15
    r_ack = p_ack.add_run("The author acknowledges the institutional support provided by the National University of Battambang (NUBB) and the computational infrastructure provided for single-GPU experimental benchmarking.")
    r_ack.font.name = "Times New Roman"
    r_ack.font.size = Pt(11)

    doc.save(output_path)
    print(f"Generated EAIT Title Page: {output_path} ({os.path.getsize(output_path)} bytes)")

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    title_path = os.path.join(base_dir, 'papers', 'Title_Page_EAIT.docx')
    create_eait_title_page(title_path)
    # Also sync root
    root_title_path = os.path.join(os.path.dirname(base_dir), 'papers', 'Title_Page_EAIT.docx')
    os.makedirs(os.path.dirname(root_title_path), exist_ok=True)
    import shutil
    shutil.copyfile(title_path, root_title_path)
