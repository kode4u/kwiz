import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_math_runs(paragraph, math_str, base_font_size=11, is_bold=False, text_color=RGBColor(0x1a, 0x1a, 0x1a)):
    """
    Renders LaTeX math expressions into clean Word text runs with proper
    italics, subscripts (e.g. T_E2E -> T with subscript E2E), superscripts,
    and Unicode mathematical symbols.
    """
    s = math_str.strip()
    
    # 1. Delimiters (must run BEFORE \le so \left is not mangled into ≤ft)
    s = s.replace(r'\left(', '(')
    s = s.replace(r'\right)', ')')
    s = s.replace(r'\left[', '[')
    s = s.replace(r'\right]', ']')
    s = s.replace(r'\left\{', '{')
    s = s.replace(r'\right\}', '}')
    s = s.replace(r'\left|', '|')
    s = s.replace(r'\right|', '|')
    
    # 2. Relational & Mathematical symbols
    s = re.sub(r'\\le(?![a-zA-Z])', '≤', s)
    s = re.sub(r'\\ge(?![a-zA-Z])', '≥', s)
    s = re.sub(r'\\in(?![a-zA-Z])', '∈', s)
    s = s.replace(r'\times', '×')
    s = s.replace(r'\pm', '±')
    s = s.replace(r'\approx', '≈')
    s = s.replace(r'\kappa', 'κ')
    s = s.replace(r'\Delta', 'Δ')
    s = s.replace(r'\dots', '…')
    s = s.replace(r'\cdots', '…')
    s = s.replace(r'\|', '‖')
    s = s.replace(r'\cdot', '·')
    s = s.replace(r'\,', ' ')
    s = s.replace(r'\quad', '   ')
    s = s.replace(r'\qquad', '     ')
    s = s.replace(r'\{', '{')
    s = s.replace(r'\}', '}')
    s = s.replace(r'\%', '%')
    
    # 3. Balanced fraction replacement: \frac{A}{B} -> (A / B)
    while r'\frac{' in s:
        idx = s.find(r'\frac{')
        depth = 0
        num_end = -1
        for j in range(idx + 5, len(s)):
            if s[j] == '{':
                depth += 1
            elif s[j] == '}':
                if depth == 0:
                    num_end = j
                    break
                depth -= 1
        if num_end != -1 and num_end + 1 < len(s) and s[num_end + 1] == '{':
            depth = 0
            den_end = -1
            for k in range(num_end + 2, len(s)):
                if s[k] == '{':
                    depth += 1
                elif s[k] == '}':
                    if depth == 0:
                        den_end = k
                        break
                    depth -= 1
            if den_end != -1:
                num = s[idx + 6:num_end]
                den = s[num_end + 2:den_end]
                s = s[:idx] + f"({num} / {den})" + s[den_end + 1:]
            else:
                break
        else:
            break
            
    # Tokenizer pattern for subscripts, superscripts, bold, text
    token_pattern = (
        r'([A-Za-z0-9]+)_\{(?:\\text\{)?([A-Za-z0-9_\-]+)\}+|'
        r'([A-Za-z0-9]+)_([A-Za-z0-9])|'
        r'([A-Za-z0-9]+)\^\{(?:\\text\{)?([A-Za-z0-9_\-]+)\}+|'
        r'([A-Za-z0-9]+)\^([A-Za-z0-9])|'
        r'\\mathbf\{([^}]+)\}|'
        r'\\text\{([^}]+)\}'
    )
    
    math_vars = {'T', 'K', 'M', 'N', 'C', 'p', 'q', 'd', 't', 'x', 'i', 'y', 'z', 'k'}
    
    pos = 0
    for match in re.finditer(token_pattern, s):
        start, end = match.span()
        if start > pos:
            plain = s[pos:start]
            sub_tokens = re.split(r'([A-Za-z]+|[^A-Za-z]+)', plain)
            for st in sub_tokens:
                if not st:
                    continue
                r = paragraph.add_run(st)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(base_font_size)
                r.font.color.rgb = text_color
                if is_bold:
                    r.bold = True
                if st in math_vars:
                    r.italic = True
                    
        g = match.groups()
        if g[0] and g[1]: # X_{sub}
            base = g[0]
            sub = g[1].replace(r'\text{', '').replace('}', '')
            r_base = paragraph.add_run(base)
            r_base.font.name = 'Times New Roman'
            r_base.font.size = Pt(base_font_size)
            r_base.font.color.rgb = text_color
            if is_bold:
                r_base.bold = True
            if base in math_vars or len(base) == 1:
                r_base.italic = True
                
            r_sub = paragraph.add_run(sub)
            r_sub.font.name = 'Times New Roman'
            r_sub.font.size = Pt(base_font_size)
            r_sub.font.color.rgb = text_color
            r_sub.font.subscript = True
            if is_bold:
                r_sub.bold = True
        elif g[2] and g[3]: # X_i
            base = g[2]
            sub = g[3]
            r_base = paragraph.add_run(base)
            r_base.font.name = 'Times New Roman'
            r_base.font.size = Pt(base_font_size)
            r_base.font.color.rgb = text_color
            if is_bold:
                r_base.bold = True
            if base in math_vars or len(base) == 1:
                r_base.italic = True
                
            r_sub = paragraph.add_run(sub)
            r_sub.font.name = 'Times New Roman'
            r_sub.font.size = Pt(base_font_size)
            r_sub.font.color.rgb = text_color
            r_sub.font.subscript = True
            if is_bold:
                r_sub.bold = True
        elif g[4] and g[5]: # X^{sup}
            base = g[4]
            sup = g[5].replace(r'\text{', '').replace('}', '')
            r_base = paragraph.add_run(base)
            r_base.font.name = 'Times New Roman'
            r_base.font.size = Pt(base_font_size)
            r_base.font.color.rgb = text_color
            if is_bold:
                r_base.bold = True
            if base in math_vars or len(base) == 1:
                r_base.italic = True
                
            r_sup = paragraph.add_run(sup)
            r_sup.font.name = 'Times New Roman'
            r_sup.font.size = Pt(base_font_size)
            r_sup.font.color.rgb = text_color
            r_sup.font.superscript = True
            if is_bold:
                r_sup.bold = True
        elif g[6] and g[7]: # X^2
            base = g[6]
            sup = g[7]
            r_base = paragraph.add_run(base)
            r_base.font.name = 'Times New Roman'
            r_base.font.size = Pt(base_font_size)
            r_base.font.color.rgb = text_color
            if is_bold:
                r_base.bold = True
            if base in math_vars or len(base) == 1:
                r_base.italic = True
                
            r_sup = paragraph.add_run(sup)
            r_sup.font.name = 'Times New Roman'
            r_sup.font.size = Pt(base_font_size)
            r_sup.font.color.rgb = text_color
            r_sup.font.superscript = True
            if is_bold:
                r_sup.bold = True
        elif g[8]: # \mathbf{x}
            r_bold = paragraph.add_run(g[8])
            r_bold.font.name = 'Times New Roman'
            r_bold.font.size = Pt(base_font_size)
            r_bold.font.color.rgb = text_color
            r_bold.bold = True
        elif g[9]: # \text{word}
            r_txt = paragraph.add_run(g[9])
            r_txt.font.name = 'Times New Roman'
            r_txt.font.size = Pt(base_font_size)
            r_txt.font.color.rgb = text_color
            if is_bold:
                r_txt.bold = True
            
        pos = end
        
    if pos < len(s):
        plain = s[pos:]
        sub_tokens = re.split(r'([A-Za-z]+|[^A-Za-z]+)', plain)
        for st in sub_tokens:
            if not st:
                continue
            r = paragraph.add_run(st)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(base_font_size)
            r.font.color.rgb = text_color
            if is_bold:
                r.bold = True
            if st in math_vars:
                r.italic = True

def add_formatted_text(paragraph, text, base_font_size=11, is_italic=False, is_bold=False, text_color=RGBColor(0x1a, 0x1a, 0x1a)):
    # Regex to tokenize bold, italic, code, and inline math
    pattern = r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\$[^$\n]+?\$)'
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            inner = part[2:-2]
            if '$' in inner:
                sub_parts = re.split(r'(\$[^$\n]+?\$)', inner)
                for sp in sub_parts:
                    if not sp:
                        continue
                    if sp.startswith('$') and sp.endswith('$'):
                        add_math_runs(paragraph, sp[1:-1], base_font_size=base_font_size, is_bold=True, text_color=text_color)
                    else:
                        r = paragraph.add_run(sp)
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(base_font_size)
                        r.font.color.rgb = text_color
                        r.bold = True
                        r.italic = is_italic
            else:
                run = paragraph.add_run(inner)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = text_color
                run.bold = True
                run.italic = is_italic
        elif part.startswith('*') and part.endswith('*'):
            inner = part[1:-1]
            if '$' in inner:
                sub_parts = re.split(r'(\$[^$\n]+?\$)', inner)
                for sp in sub_parts:
                    if not sp:
                        continue
                    if sp.startswith('$') and sp.endswith('$'):
                        add_math_runs(paragraph, sp[1:-1], base_font_size=base_font_size, is_bold=is_bold, text_color=text_color)
                    else:
                        r = paragraph.add_run(sp)
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(base_font_size)
                        r.font.color.rgb = text_color
                        r.italic = True
                        r.bold = is_bold
            else:
                run = paragraph.add_run(inner)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = text_color
                run.italic = True
                run.bold = is_bold
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(base_font_size - 0.5)
            run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
            if is_bold:
                run.bold = True
        elif part.startswith('$') and part.endswith('$'):
            add_math_runs(paragraph, part[1:-1], base_font_size=base_font_size, is_bold=is_bold, text_color=text_color)
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = text_color
            run.bold = is_bold
            run.italic = is_italic



def convert_md_to_docx(md_path, docx_path):
    doc = Document()
    
    # Page Margins: 1 inch (72 pt) all sides
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    i = 0
    n = len(lines)
    
    while i < n:
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()
        
        # Blank line
        if not stripped:
            i += 1
            continue
            
        # Title (# Title)
        if stripped.startswith('# '):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(14)
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
            i += 1
            continue
            
        # Author / Affiliation blocks
        if stripped.startswith('<div class="author-block">'):
            author_text = re.sub(r'<[^>]+>', '', stripped).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(author_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True
            i += 1
            continue
            
        if stripped.startswith('<div class="affil-block">'):
            affil_text = re.sub(r'<br\s*/?>', '\n', stripped)
            affil_text = re.sub(r'<[^>]+>', '', affil_text).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(14)
            for subline in affil_text.split('\n'):
                run = p.add_run(subline + '\n')
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                run.italic = True
                run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            i += 1
            continue
            
        if stripped == '---':
            i += 1
            continue
            
        # Heading 1 (## Heading)
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, h_text, base_font_size=14, is_bold=True, text_color=RGBColor(0x0f, 0x17, 0x2a))
            i += 1
            continue
            
        # Heading 2 (### Heading)
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, h_text, base_font_size=12, is_bold=True, text_color=RGBColor(0x1e, 0x29, 0x3b))
            i += 1
            continue

        # Heading 3 (#### Heading)
        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, h_text, base_font_size=11, is_bold=True, is_italic=True, text_color=RGBColor(0x33, 0x41, 0x55))
            i += 1
            continue

        # Image ![caption](path)
        if stripped.startswith('![') and '](' in stripped:
            img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
            if img_match:
                caption = img_match.group(1)
                img_path = img_match.group(2)
                # Map SVG to generated PNG
                fig_dir = os.path.join(os.path.dirname(md_path), 'figures')
                if 'pipeline_architecture' in img_path:
                    png_path = os.path.join(fig_dir, 'pipeline_architecture.png')
                elif 'cache_decision_flow' in img_path:
                    png_path = os.path.join(fig_dir, 'cache_decision_flow.png')
                else:
                    png_path = os.path.join(fig_dir, os.path.basename(img_path))
                
                if os.path.exists(png_path):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(8)
                    p_img.paragraph_format.space_after = Pt(2)
                    run_img = p_img.add_run()
                    run_img.add_picture(png_path, width=Inches(6.2))
                    
                i += 1
                continue

        # Figure caption (*Figure X: ...*)
        if stripped.startswith('*Figure ') and stripped.endswith('*'):
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
            add_formatted_text(p_cap, stripped[1:-1], base_font_size=9.5, is_italic=True, text_color=RGBColor(0x47, 0x55, 0x69))
            i += 1
            continue

        # Blockquote (> text)
        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_formatted_text(p, stripped[2:], base_font_size=10.5, is_italic=True)
            i += 1
            continue

        # Display Math ($$...$$)
        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4:
            math_content = stripped[2:-2].strip()
            if 'sim' in math_content:
                omml_eq1 = '''<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
  <m:oMath>
    <m:r><m:rPr><m:nor/></m:rPr><m:t>sim</m:t></m:r>
    <m:d><m:e><m:r><m:rPr><m:b/></m:rPr><m:t>q</m:t></m:r><m:r><m:t>, </m:t></m:r><m:r><m:rPr><m:b/></m:rPr><m:t>d</m:t></m:r></m:e></m:d>
    <m:r><m:t> = </m:t></m:r>
    <m:f>
      <m:num><m:r><m:rPr><m:b/></m:rPr><m:t>q</m:t></m:r><m:r><m:t> · </m:t></m:r><m:r><m:rPr><m:b/></m:rPr><m:t>d</m:t></m:r></m:num>
      <m:den>
        <m:sSub>
          <m:e><m:d><m:dPr><m:begChr m:val="‖"/><m:endChr m:val="‖"/></m:dPr><m:e><m:r><m:rPr><m:b/></m:rPr><m:t>q</m:t></m:r></m:e></m:d></m:e>
          <m:sub><m:r><m:t>2</m:t></m:r></m:sub>
        </m:sSub>
        <m:r><m:t> </m:t></m:r>
        <m:sSub>
          <m:e><m:d><m:dPr><m:begChr m:val="‖"/><m:endChr m:val="‖"/></m:dPr><m:e><m:r><m:rPr><m:b/></m:rPr><m:t>d</m:t></m:r></m:e></m:d></m:e>
          <m:sub><m:r><m:t>2</m:t></m:r></m:sub>
        </m:sSub>
      </m:den>
    </m:f>
  </m:oMath>
</m:oMathPara>'''
                p_math = doc.add_paragraph()
                p_math.paragraph_format.space_before = Pt(6)
                p_math.paragraph_format.space_after = Pt(6)
                p_math._p.append(parse_xml(omml_eq1))
            elif 'CacheHitRate' in math_content:
                omml_eq2 = '''<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
  <m:oMath>
    <m:r><m:rPr><m:nor/></m:rPr><m:t>CacheHitRate</m:t></m:r>
    <m:r><m:t> = </m:t></m:r>
    <m:d>
      <m:e>
        <m:f>
          <m:num><m:r><m:rPr><m:nor/></m:rPr><m:t>ReusedChunks</m:t></m:r></m:num>
          <m:den><m:r><m:rPr><m:nor/></m:rPr><m:t>EligibleChunks</m:t></m:r></m:den>
        </m:f>
      </m:e>
    </m:d>
    <m:r><m:t> × 100%,     </m:t></m:r>
    <m:r><m:rPr><m:nor/></m:rPr><m:t>Speedup</m:t></m:r>
    <m:r><m:t> = </m:t></m:r>
    <m:f>
      <m:num><m:sSub><m:e><m:r><m:t>T</m:t></m:r></m:e><m:sub><m:r><m:rPr><m:nor/></m:rPr><m:t>full</m:t></m:r></m:sub></m:sSub></m:num>
      <m:den><m:sSub><m:e><m:r><m:t>T</m:t></m:r></m:e><m:sub><m:r><m:rPr><m:nor/></m:rPr><m:t>incremental</m:t></m:r></m:sub></m:sSub></m:den>
    </m:f>
  </m:oMath>
</m:oMathPara>'''
                p_math = doc.add_paragraph()
                p_math.paragraph_format.space_before = Pt(6)
                p_math.paragraph_format.space_after = Pt(6)
                p_math._p.append(parse_xml(omml_eq2))
            elif 'Throughput' in math_content:
                omml_eq3 = '''<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
  <m:oMath>
    <m:r><m:rPr><m:nor/></m:rPr><m:t>Throughput</m:t></m:r>
    <m:r><m:t> = </m:t></m:r>
    <m:f>
      <m:num><m:sSub><m:e><m:r><m:t>N</m:t></m:r></m:e><m:sub><m:r><m:rPr><m:nor/></m:rPr><m:t>successful</m:t></m:r></m:sub></m:sSub></m:num>
      <m:den><m:r><m:t>Δ</m:t></m:r><m:r><m:t>t</m:t></m:r></m:den>
    </m:f>
    <m:r><m:t>,     </m:t></m:r>
    <m:r><m:rPr><m:nor/></m:rPr><m:t>SuccessRate</m:t></m:r>
    <m:r><m:t> = </m:t></m:r>
    <m:d>
      <m:e>
        <m:f>
          <m:num><m:sSub><m:e><m:r><m:t>N</m:t></m:r></m:e><m:sub><m:r><m:rPr><m:nor/></m:rPr><m:t>successful</m:t></m:r></m:sub></m:sSub></m:num>
          <m:den><m:sSub><m:e><m:r><m:t>N</m:t></m:r></m:e><m:sub><m:r><m:rPr><m:nor/></m:rPr><m:t>total</m:t></m:r></m:sub></m:sSub></m:den>
        </m:f>
      </m:e>
    </m:d>
    <m:r><m:t> × 100%</m:t></m:r>
  </m:oMath>
</m:oMathPara>'''
                p_math = doc.add_paragraph()
                p_math.paragraph_format.space_before = Pt(6)
                p_math.paragraph_format.space_after = Pt(6)
                p_math._p.append(parse_xml(omml_eq3))
            else:
                p_math = doc.add_paragraph()
                p_math.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_math.paragraph_format.space_before = Pt(6)
                p_math.paragraph_format.space_after = Pt(6)
                add_math_runs(p_math, math_content, base_font_size=11)
            i += 1
            continue

        # Tables (| col1 | col2 |)
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
                
            if len(table_lines) >= 2:
                # First line: headers
                headers = [c.strip() for c in table_lines[0][1:-1].split('|')]
                # Second line: alignment / separator
                # Data lines:
                data_rows = []
                for tl in table_lines[2:]:
                    cols = [c.strip() for c in tl[1:-1].split('|')]
                    data_rows.append(cols)
                    
                table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True
                
                # Format header row
                hdr_cells = table.rows[0].cells
                for col_idx, h_text in enumerate(headers):
                    hdr_cells[col_idx].text = ""
                    p = hdr_cells[col_idx].paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                    add_formatted_text(p, h_text, base_font_size=9.5, is_bold=True)
                    set_cell_background(hdr_cells[col_idx], "F1F5F9")
                    set_cell_margins(hdr_cells[col_idx], top=80, bottom=80, left=120, right=120)
                    
                # Format data rows
                for r_idx, row_data in enumerate(data_rows):
                    row_cells = table.rows[r_idx + 1].cells
                    bg_color = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
                    for col_idx, cell_value in enumerate(row_data):
                        if col_idx < len(row_cells):
                            row_cells[col_idx].text = ""
                            p = row_cells[col_idx].paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            # Align center for numbers, left for labels
                            is_center = col_idx > 0 or cell_value.isdigit()
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_center else WD_ALIGN_PARAGRAPH.LEFT
                            add_formatted_text(p, cell_value, base_font_size=9)
                            set_cell_background(row_cells[col_idx], bg_color)
                            set_cell_margins(row_cells[col_idx], top=60, bottom=60, left=100, right=100)
                            
                # Add spacing after table
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_before = Pt(0)
                p_after.paragraph_format.space_after = Pt(6)
            continue

        # Bullet list (- or * or •)
        if re.match(r'^[-*•]\s+', stripped):
            item_text = re.sub(r'^[-*•]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, item_text, base_font_size=11)
            i += 1
            continue

        # Numbered list (1. 2. etc)
        if re.match(r'^\d+\.\s+', stripped):
            item_text = re.sub(r'^\d+\.\s+', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, item_text, base_font_size=11)
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        # Check if it is the References section
        if stripped.startswith('[') and ']' in stripped and re.match(r'^\[\d+\]', stripped):
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            add_formatted_text(p, stripped, base_font_size=9.5)
        else:
            add_formatted_text(p, stripped, base_font_size=11)
            
        i += 1

    doc.save(docx_path)
    print(f"Generated {docx_path} ({os.path.getsize(docx_path)} bytes)")

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    md_file = os.path.join(base_dir, 'papers', 'paper.md')
    docx_file = os.path.join(base_dir, 'papers', 'paper.docx')
    convert_md_to_docx(md_file, docx_file)
